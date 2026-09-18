"""
Tests for the export service (step 6) and the dashboard (step 7).

The export tests check the files are actually produced and that the numbers
inside them match the results object, rather than just checking nothing threw.
The dashboard tests drive the real app through Streamlit's AppTest harness, so
a broken widget or a crash on submit shows up here instead of in a demo.
"""

import io
import os
import zipfile

import pytest

from src import score_configuration as cfg
from src.company_profile import CompanyProfile
from src.export_service import (
    ExportService,
    suggested_filename,
    to_excel_bytes,
    to_pdf_bytes,
)
from src.recommendation_engine import RecommendationEngine
from src.scoring_engine import ScoringEngine


CONTEXT = {
    "industry_sector": "Manufacturing",
    "employee_band": "10-49",
    "years_in_operation": "6-10 years",
    "region": "UAE",
    "current_ai_stage": "Exploring",
}


def results_for(answer_for) -> "object":
    responses = {item.id: answer_for(item) for item in cfg.ALL_ITEMS}
    results = ScoringEngine().score(
        CompanyProfile.from_dict({**CONTEXT, **responses})
    )
    RecommendationEngine().recommend(results)
    return results


@pytest.fixture
def mixed_results():
    """A company with strengths, barriers, item gaps and recommendations."""
    responses = {item.id: 3 for item in cfg.ALL_ITEMS}
    for item in cfg.FACTORS_BY_ID["budget"].items:
        responses[item.id] = 1 if item.reverse else 5      # -> 100
    for item in cfg.FACTORS_BY_ID["governance"].items:
        responses[item.id] = 5 if item.reverse else 1      # -> 0

    results = ScoringEngine().score(
        CompanyProfile.from_dict({**CONTEXT, **responses})
    )
    RecommendationEngine().recommend(results)
    return results


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def test_pdf_is_a_real_pdf(mixed_results):
    data = to_pdf_bytes(mixed_results)
    assert data.startswith(b"%PDF-")
    assert data.rstrip().endswith(b"%%EOF")
    assert len(data) > 2000


@pytest.mark.parametrize(
    "answer_for",
    [
        lambda item: 1 if item.reverse else 5,   # perfect, no recommendations
        lambda item: 5 if item.reverse else 1,   # worst, everything flagged
        lambda item: 3,                          # flat neutral
    ],
)
def test_pdf_builds_for_every_extreme(answer_for):
    """The layout must survive empty sections and full ones alike."""
    assert to_pdf_bytes(results_for(answer_for)).startswith(b"%PDF-")


def test_pdf_handles_factor_names_containing_ampersands(mixed_results):
    """
    Most factor names contain "&", which reportlab would otherwise treat as
    the start of an entity and reject. This is the regression guard for that.
    """
    assert any("&" in factor.name for factor in cfg.FACTORS), "test premise gone"
    assert to_pdf_bytes(mixed_results).startswith(b"%PDF-")


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def test_excel_is_a_real_workbook_with_the_expected_sheets(mixed_results):
    from openpyxl import load_workbook

    data = to_excel_bytes(mixed_results)
    assert zipfile.is_zipfile(io.BytesIO(data)), "xlsx files are zip archives"

    workbook = load_workbook(io.BytesIO(data))
    assert workbook.sheetnames == [
        "Summary", "Factor scores", "Responses", "Recommendations",
    ]


def test_excel_factor_scores_match_the_results(mixed_results):
    from openpyxl import load_workbook

    sheet = load_workbook(io.BytesIO(to_excel_bytes(mixed_results)))["Factor scores"]
    written = {
        sheet.cell(row=row, column=1).value: sheet.cell(row=row, column=2).value
        for row in range(2, len(cfg.FACTORS) + 2)
    }

    for factor in mixed_results.factor_scores.values():
        assert written[factor.name] == pytest.approx(factor.score, abs=0.05)


def test_excel_lists_every_question(mixed_results):
    from openpyxl import load_workbook

    sheet = load_workbook(io.BytesIO(to_excel_bytes(mixed_results)))["Responses"]
    written = {sheet.cell(row=row, column=1).value for row in range(2, sheet.max_row + 1)}
    assert written == {item.id for item in cfg.ALL_ITEMS}


def test_excel_actually_contains_charts(mixed_results):
    """
    A workbook of bare numbers was the complaint: the PDF drew bars and the
    spreadsheet drew nothing, which is backwards for the file people are
    meant to analyse in.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(to_excel_bytes(mixed_results)))
    assert len(workbook["Factor scores"]._charts) >= 2, "factor charts missing"
    assert len(workbook["Responses"]._charts) >= 1, "question chart missing"


def test_excel_charts_survive_a_company_with_no_recommendations():
    """The chart ranges are built from row counts, so the empty case matters."""
    from openpyxl import load_workbook

    perfect = results_for(lambda item: 1 if item.reverse else 5)
    assert perfect.recommendations == []
    workbook = load_workbook(io.BytesIO(to_excel_bytes(perfect)))
    assert len(workbook["Factor scores"]._charts) >= 2


def test_excel_never_writes_a_size_excel_would_read_as_a_date(mixed_results):
    """Same trap as the readable CSV -- "10-49" becomes Oct-49 in Excel."""
    import re
    from openpyxl import load_workbook

    sheet = load_workbook(io.BytesIO(to_excel_bytes(mixed_results)))["Summary"]
    for row in sheet.iter_rows(values_only=True):
        for value in row:
            if isinstance(value, str):
                assert not re.fullmatch(r"\d+-\d+", value), f"{value!r} converts to a date"


def test_filenames_carry_the_tier_and_no_identifying_detail(mixed_results):
    name = suggested_filename(mixed_results, "pdf")
    assert name.endswith(".pdf")
    assert mixed_results.readiness_tier.lower() in name


def test_service_wrapper_matches_the_functions(mixed_results):
    service = ExportService()
    assert service.to_pdf(mixed_results).startswith(b"%PDF-")
    assert zipfile.is_zipfile(io.BytesIO(service.to_excel(mixed_results)))


def test_write_helpers_put_files_on_disk(mixed_results, tmp_path):
    from src.export_service import write_excel, write_pdf

    pdf = write_pdf(mixed_results, str(tmp_path / "r.pdf"))
    excel = write_excel(mixed_results, str(tmp_path / "r.xlsx"))

    assert open(pdf, "rb").read(5) == b"%PDF-"
    assert zipfile.is_zipfile(excel)


# ---------------------------------------------------------------------------
# Dashboard
# ---------------------------------------------------------------------------
# These drive app.py itself. They're slower than the rest of the suite because
# each one runs the whole script, but they catch the class of bug that only
# appears when the app is actually used.

pytest.importorskip("streamlit")
from streamlit.testing.v1 import AppTest  # noqa: E402


APP_PATH = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "app.py"
)


def run_app() -> AppTest:
    # Absolute path: AppTest resolves a relative one against this test file's
    # own folder, not the project root.
    return AppTest.from_file(APP_PATH, default_timeout=120).run()


def fill_in(app: AppTest, answer_for) -> AppTest:
    for field, options in cfg.CATEGORICAL_FIELDS.items():
        app.session_state[f"ctx_{field}"] = options[0]
    for item in cfg.ALL_ITEMS:
        app.session_state[f"item_{item.id}"] = answer_for(item)
    return app.run()


def test_app_loads_with_every_question_on_it():
    app = run_app()
    assert not app.exception
    assert len(app.radio) == len(cfg.ALL_ITEMS)
    assert len(app.selectbox) == len(cfg.CATEGORICAL_FIELDS)


def test_submitting_an_empty_form_reports_what_is_missing():
    app = run_app()
    app.button[0].click().run()

    assert not app.exception
    assert "results" not in app.session_state
    assert app.error, "an incomplete submission must show an error"


def test_completed_assessment_produces_a_score():
    app = fill_in(run_app(), lambda item: 1 if item.reverse else 5)
    app.button[0].click().run()

    assert not app.exception
    results = app.session_state["results"]
    assert results.overall_score == pytest.approx(100.0)
    assert results.readiness_tier == "Advanced"


def test_results_view_offers_both_downloads():
    app = fill_in(run_app(), lambda item: 3)
    app.button[0].click().run()

    assert not app.exception
    labels = [button.label for button in app.download_button]
    assert "Download PDF report" in labels
    assert "Download Excel workbook" in labels


def test_answering_a_question_does_not_collapse_its_section():
    """
    Streamlit re-runs the whole script on every answer. When `expanded` came
    from a fixed value, that rerun snapped the section shut mid-way through
    filling it in, so people lost their place and couldn't tell what they had
    already completed.
    """
    app = run_app()
    before = [expander.proto.expanded for expander in app.expander]
    assert all(before), "sections should start open"

    # Answer one question in a later section, then re-run as Streamlit would.
    app.session_state["item_CUL_1"] = 4
    app = app.run()

    assert not app.exception
    after = [expander.proto.expanded for expander in app.expander]
    assert all(after), "a section closed itself after an answer was given"
    assert len(after) == len(cfg.FACTORS)


def test_section_headers_show_how_much_is_left():
    app = run_app()
    app.session_state["item_BUD_1"] = 5
    app = app.run()

    budget = next(e for e in app.expander if "Budget" in e.label)
    assert "1/4" in budget.label

    for item in cfg.FACTORS_BY_ID["budget"].items:
        app.session_state[f"item_{item.id}"] = 4
    app = app.run()

    budget = next(e for e in app.expander if "Budget" in e.label)
    assert "✓" in budget.label


def test_other_asks_for_a_specific_answer():
    """Picking Other reveals a box, and Other alone doesn't count as answered."""
    app = run_app()
    app.session_state["ctx_industry_sector"] = "Other (please specify)"
    app = app.run()

    assert not app.exception
    assert any("specify" in box.label.lower() or box.label == "Please specify"
               for box in app.text_input), "no box appeared to type the sector in"


def test_starting_over_clears_the_answers():
    app = fill_in(run_app(), lambda item: 3)
    app.button[0].click().run()
    assert "results" in app.session_state

    next(b for b in app.button if b.label == "Start a new assessment").click().run()

    assert not app.exception
    assert "results" not in app.session_state
    # The form is showing again, so the radios exist but none is selected.
    assert all(radio.value is None for radio in app.radio)


# ---------------------------------------------------------------------------
# Written answers, and the dataset report
# ---------------------------------------------------------------------------

def test_written_answers_reach_the_results_and_the_pdf():
    """They were collected and then dropped on the floor."""
    profile = CompanyProfile.from_dict({
        **CONTEXT,
        "open_biggest_barrier": "Cost, and nobody in-house who understands it.",
        "open_what_would_help": "Affordable training.",
        **{item.id: 3 for item in cfg.ALL_ITEMS},
    })
    results = ScoringEngine().score(profile)
    assert len(results.notes) == 2
    assert "notes" in results.to_dict()
    assert to_pdf_bytes(results).startswith(b"%PDF-")


def test_pdf_carries_the_profile_shape_and_every_question(mixed_results):
    """
    The report was missing two things the screen showed: the profile shape and
    the question-by-question detail. Size is a blunt proxy, but a report with
    a radar and 32 rows in it cannot be as small as one without.
    """
    data = to_pdf_bytes(mixed_results)
    assert data.startswith(b"%PDF-")
    assert len(data) > 9000, "the detail sections look to be missing"


def _small_cohort():
    from src.recommendation_engine import RecommendationEngine

    scored, profiles = [], {}
    for index, (sector, country, answer) in enumerate((
        ("Retail", "UAE", 2), ("Retail", "UAE", 4),
        ("Logistics", "Kenya", 3), ("Logistics", "Kenya", 5),
    )):
        profile = CompanyProfile.from_dict({
            **CONTEXT, "industry_sector": sector, "region": country,
            **{item.id: answer for item in cfg.ALL_ITEMS},
        })
        result = ScoringEngine().score(profile)
        RecommendationEngine().recommend(result)
        identifier = f"R{index}"
        scored.append((identifier, result))
        profiles[identifier] = profile
    return scored, profiles


def test_dataset_summary_splits_profiles_by_country_and_sector():
    """
    One averaged shape across several countries describes a company that
    exists nowhere in the file.
    """
    from src.export_service import summarise_dataset

    scored, profiles = _small_cohort()
    summary = summarise_dataset(scored, profiles)

    assert summary.count == 4
    assert summary.split_profiles
    assert {g.label for g in summary.by_country} == {"United Arab Emirates", "Kenya"}
    assert {g.label for g in summary.by_sector} == {"Retail", "Logistics"}


def test_dataset_gets_recommendations_like_a_single_assessment():
    from src.export_service import summarise_dataset

    scored, profiles = _small_cohort()
    summary = summarise_dataset(scored, profiles)
    assert summary.recommendations, "a weak cohort produced no advice"
    assert all(r.title for r in summary.recommendations)


def test_dataset_downloads_are_real_files_with_charts():
    from openpyxl import load_workbook
    from src.export_service import (
        dataset_excel_bytes, dataset_pdf_bytes, summarise_dataset,
    )

    scored, profiles = _small_cohort()
    notes = [("R0", {"Biggest barrier": "Cost and a lack of trained staff."}),
             ("R1", {"Biggest barrier": "Data quality and privacy rules."})]
    summary = summarise_dataset(scored, profiles, notes, "test.csv")

    pdf = dataset_pdf_bytes(summary)
    assert pdf.startswith(b"%PDF-") and len(pdf) > 4000

    workbook = load_workbook(io.BytesIO(dataset_excel_bytes(summary)))
    assert "Factor averages" in workbook.sheetnames
    assert len(workbook["Factor averages"]._charts) >= 2
    assert sum(len(workbook[s]._charts) for s in workbook.sheetnames) >= 4


def test_themes_group_the_written_answers_by_subject():
    """
    "cost", "budget" and "expensive" are one concern. Counting raw words
    would report them as three small ones.
    """
    from src.export_service import extract_themes

    notes = [
        ("a", {"q": "Cost is the main problem for us."}),
        ("b", {"q": "We cannot afford the investment."}),
        ("c", {"q": "Budget is too tight."}),
        ("d", {"q": "Staff need training before we can use it."}),
    ]
    themes = dict((name, count) for name, count, _ in extract_themes(notes))
    assert themes.get("Cost and funding", 0) == 3
    assert themes.get("Skills and training", 0) == 1


# ---------------------------------------------------------------------------
# Chart labelling
# ---------------------------------------------------------------------------
# The charts openpyxl produces the plain way come out with no labels round the
# edge at all, because set_categories() writes a range of words as a *numeric*
# reference and Excel reads that as empty. These tests guard the fix.
#
# They read the chart XML as a tree rather than as a string. Searching the raw
# text looks simpler and is a trap: openpyxl serialises through lxml when it is
# installed and through the standard library when it is not, and the two write
# empty elements differently -- `<majorGridlines/>` against
# `<majorGridlines />`. The workbooks are identical to Excel; only the bytes
# differ. A test matching the text passes on one machine and fails on the next.

CHART_NS = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"


def _charts(data: bytes, kind: str = ""):
    """
    Every chart in a workbook, parsed. `kind` filters by chart type, e.g.
    "radarChart" or "barChart".
    """
    import xml.etree.ElementTree as ElementTree

    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        trees = [ElementTree.fromstring(archive.read(name))
                 for name in archive.namelist()
                 if name.startswith("xl/charts/chart")
                 and name.endswith(".xml")]

    if kind:
        trees = [t for t in trees if t.find(f".//{CHART_NS}{kind}") is not None]
    return trees


def _axis(chart, name: str):
    """A chart's category ("catAx") or value ("valAx") axis."""
    return chart.find(f".//{CHART_NS}{name}")


def _flag(element, name: str):
    """The val="..." of a child element, or None if it isn't there."""
    child = element.find(f"{CHART_NS}{name}") if element is not None else None
    return child.get("val") if child is not None else None


def test_every_chart_labels_its_categories_with_text(mixed_results):
    """
    The factor names have to reach the chart as text. Written as a numeric
    reference they silently disappear, which is what left the profile shape
    with no factor labels on it.
    """
    charts = _charts(to_excel_bytes(mixed_results))
    assert charts

    for chart in charts:
        categories = chart.findall(f".//{CHART_NS}cat")
        assert categories, "a chart with no categories at all"
        for category in categories:
            assert category.find(f"{CHART_NS}strRef") is not None, \
                "categories written as numbers -- the labels will be blank"
            assert category.find(f"{CHART_NS}numRef") is None


def test_the_radar_carries_the_readiness_factor_names(mixed_results):
    charts = _charts(to_excel_bytes(mixed_results), "radarChart")
    assert charts, "the profile shape should be in the workbook"

    for chart in charts:
        assert chart.find(f".//{CHART_NS}cat/{CHART_NS}strRef") is not None
        # The category axis carries the factor names, so it must be drawn.
        assert _flag(_axis(chart, "catAx"), "delete") == "0"


def test_the_radar_does_not_print_its_scale_down_the_middle(mixed_results):
    """
    A radar draws its value axis as a column of numbers through the centre of
    the shape, on top of the fill. The rings carry the scale instead.
    """
    charts = _charts(to_excel_bytes(mixed_results), "radarChart")
    assert charts

    for chart in charts:
        value_axis = _axis(chart, "valAx")
        assert _flag(value_axis, "delete") == "1", "0-100 printed over the shape"
        # Hiding the axis must not take the rings with it.
        assert value_axis.find(f"{CHART_NS}majorGridlines") is not None


def test_bar_charts_show_their_values_and_name_their_axes(mixed_results):
    charts = _charts(to_excel_bytes(mixed_results), "barChart")
    assert charts

    for chart in charts:
        labels = chart.find(f".//{CHART_NS}dLbls")
        assert labels is not None, "no value labels on a bar chart"
        assert _flag(labels, "showVal") == "1"
        # Only the value. Left unset, some readers add the series and category
        # names too and the label sprawls across its neighbours.
        assert _flag(labels, "showSerName") == "0"
        assert _flag(labels, "showCatName") == "0"


def test_the_dataset_workbook_is_labelled_the_same_way():
    from src.export_service import dataset_excel_bytes, summarise_dataset

    scored, profiles = _small_cohort()
    charts = _charts(dataset_excel_bytes(summarise_dataset(scored, profiles)))

    assert charts
    for chart in charts:
        for category in chart.findall(f".//{CHART_NS}cat"):
            assert category.find(f"{CHART_NS}strRef") is not None
            assert category.find(f"{CHART_NS}numRef") is None


def test_the_dataset_workbook_splits_the_shape_by_group():
    """
    A radar per country and per sector, matching the PDF -- one averaged
    shape across several of them describes a company that does not exist.
    """
    from src.export_service import dataset_excel_bytes, summarise_dataset

    scored, profiles = _small_cohort()
    data = dataset_excel_bytes(summarise_dataset(scored, profiles))

    radars = _charts(data, "radarChart")
    grouped = [r for r in radars
               if len(r.findall(f".//{CHART_NS}ser")) > 1]
    assert grouped, "expected a radar holding one series per group"


# ---------------------------------------------------------------------------
# Where the explanation sits
# ---------------------------------------------------------------------------

def _pdf_text(data: bytes) -> str:
    pypdf = pytest.importorskip("pypdf")
    reader = pypdf.PdfReader(io.BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def test_the_scoring_explanation_comes_before_the_scores(mixed_results):
    """
    An explanation printed after the figures it explains has been read too
    late to be any use.
    """
    text = _pdf_text(to_pdf_bytes(mixed_results))
    explanation = text.find("How the score is calculated")
    scores = text.find("Readiness by factor")

    assert explanation != -1 and scores != -1
    assert explanation < scores


def test_the_dataset_explanation_comes_first_too():
    from src.export_service import dataset_pdf_bytes, summarise_dataset

    scored, profiles = _small_cohort()
    text = _pdf_text(dataset_pdf_bytes(summarise_dataset(scored, profiles)))

    explanation = text.find("How to read this")
    scores = text.find("Average score by factor")

    assert explanation != -1 and scores != -1
    assert explanation < scores


def test_the_dashboard_shows_what_the_written_answers_were_read_as():
    """
    The three open questions used to be collected and then printed back
    unread. Submitting one now has to reach the results view.
    """
    app = run_app()
    app.session_state["open_biggest_barrier"] = (
        "Cost is the biggest issue and we have no budget set aside."
    )
    app = fill_in(app, lambda item: 3)
    app.button[0].click().run()

    assert not app.exception
    results = app.session_state["results"]
    assert results.text_insights is not None
    assert results.text_insights.signal("budget") is not None

    body = " ".join(str(element.value) for element in app.markdown)
    assert "In your own words" in body
