"""
Survey import

Takes a spreadsheet someone actually exported -- from Google Forms, from
Excel, from wherever -- and works out which column is which, instead of
demanding the file already use the tool's internal names.

The first real survey run was the reason this exists. A Google Forms export
of the tool's own questionnaire was rejected wholesale: Forms writes the
column as "Industry Sector" where the tool wanted `industry_sector`, writes
the full question text where the tool wanted `BUD_1`, and writes "Agree"
where the tool wanted 4. Nothing about the responses was wrong. Only the
formatting was, and a tool that can't read its own questionnaire's export is
not much use to anyone.

How the matching works:

  Columns named with an item code (BUD_1) are taken as-is.

  Otherwise the question text is pulled out of the header -- Forms wraps it
  in brackets after the section name, like
  "Data Readiness [Our company data is well-organized...]" -- and compared
  against the wording in the configuration. The comparison is fuzzy on
  purpose: the live form has typos ("idenitfied", "exisitng") that the
  configuration spells correctly, and an exact match would throw those two
  questions away for no good reason.

  Context columns are matched against a list of the ways people label them.

  Answers go through parse_rating, so words, numbers, or "4 - Agree" all
  read the same.

Anything that doesn't match is reported rather than silently dropped, so it
is always possible to see what the tool did with a file.
"""

from __future__ import annotations

import csv
import io
import math
import re
from dataclasses import dataclass, field
from difflib import SequenceMatcher
from typing import Dict, List, Optional, Sequence, Tuple

from . import score_configuration as cfg


# A header has to be at least this similar to a question before it counts as
# that question. High enough that two different questions never collide, low
# enough to survive a handful of typos.
TEXT_MATCH_THRESHOLD = 0.82

# How close two questions must be in MEANING before the importer will suggest
# that one stands in for the other. Lower than the wording threshold, because
# a paraphrase never scores as high as a near-identical string, but high
# enough that unrelated questions don't get paired up. Matches at this level
# are marked as such and shown for confirmation rather than applied silently:
# deciding that another survey's question measures the same thing as one of
# ours is a research judgement, not a formatting fix.
MEANING_MATCH_THRESHOLD = 0.30

# Columns that carry no assessment data and should not be reported as
# unmatched, because nobody expects the tool to do anything with them.
IGNORED_HEADERS = {
    "timestamp", "time", "date", "submitted at", "submission time",
    "email address", "email", "id", "response id", "score", "username",
    "name", "company", "company name", "organisation", "organization",
}

# The ways the five context questions tend to get labelled.
CONTEXT_ALIASES: Dict[str, Tuple[str, ...]] = {
    "industry_sector": (
        "industry_sector", "industry sector", "industry", "sector",
        "business sector", "what sector does your company operate in",
        "which sector", "industry type", "field",
    ),
    "employee_band": (
        "employee_band", "number of employees", "no of employees", "employees",
        "company size", "size of company", "organisation size", "headcount",
        "how many people does your company employ", "employee band",
        "staff count", "number of staff",
    ),
    "years_in_operation": (
        "years_in_operation", "years in operation", "years operating",
        "company age", "age of company", "how long has the company been operating",
        "years in business", "time in operation",
    ),
    "region": (
        "region", "country", "region/country of operation",
        "region / country of operation", "country of operation",
        "location", "where does the company mainly operate", "market",
        "q1. country selection", "country selection",
    ),
    "current_ai_stage": (
        "current_ai_stage", "current stage of ai adoption", "ai adoption stage",
        "stage of ai adoption", "ai stage", "current ai stage",
        "where is the company with ai at the moment", "adoption stage",
    ),
}


# ---------------------------------------------------------------------------
# Text tidying
# ---------------------------------------------------------------------------

def _normalise(text: str) -> str:
    """Lowercase, strip punctuation, collapse whitespace."""
    cleaned = re.sub(r"[^a-z0-9 ]+", " ", str(text).lower())
    return " ".join(cleaned.split())


# ---------------------------------------------------------------------------
# Reading questions by meaning, not just by spelling
# ---------------------------------------------------------------------------
# String similarity handles the same questionnaire typed slightly differently.
# It cannot handle the same question asked in different words: "Our agency has
# sufficient and reliable IT infrastructure" and "Current IT infrastructure
# could support new AI tools" share almost no characters in that order, and a
# character-level comparison scores them near zero.
#
# What follows is a small information-retrieval pipeline: drop the words that
# carry no meaning, reduce the rest to stems, add the concept each domain word
# belongs to, weight everything by how rare it is across the 32 questions, and
# compare the resulting vectors by cosine similarity.
#
# Deliberately not a pretrained language model. A sentence-transformer would
# match a little better and would cost a large dependency, a download at
# startup, and results that can't be explained to an examiner or reproduced
# exactly a year from now. Everything here is deterministic and can be read
# off the page, which matters more for a dissertation than the last few points
# of accuracy.

STOPWORDS: frozenset = frozenset("""
a an the and or but if of in on at to for from with without by as is are was
were be been being do does did have has had having we our us you your they
their them it its this that these those there here which who whom what when
where how why can could would should will shall may might must not no nor
than then so such very more most much many few own same other another each
every any all both either neither about into over under again further once
""".split())

# Words that mean the same thing for this instrument. Each entry maps a
# concept to the vocabulary that signals it, so a question phrased in one
# organisation's language still lands on the right readiness factor.
CONCEPT_LEXICON: Dict[str, Tuple[str, ...]] = {
    "infra": (
        "infrastructure", "broadband", "server", "cloud", "hardware",
        "network", "platform", "computing", "connectivity", "bandwidth",
        "environment", "architecture", "integration", "integrated", "legacy",
        "outdated", "compatible", "incompatible",
    ),
    "money": (
        "budget", "fund", "funding", "funded", "financial", "finance",
        "cost", "costly", "money", "invest", "investment", "grant",
        "afford", "affordable", "spend", "spending", "capital", "roi",
        "return", "expenditure", "resourced",
    ),
    "people": (
        "skill", "skilled", "training", "train", "trained", "competence",
        "competency", "capability", "expertise", "knowledge", "literacy",
        "upskill", "reskill", "learn", "learning", "talent", "staff",
        "employee", "workforce", "personnel", "recruit", "hire",
    ),
    "leader": (
        "leadership", "leader", "management", "executive", "senior",
        "director", "sponsor", "champion", "champions", "strategy",
        "strategic", "vision", "priority", "prioritise", "direction",
        "mandate", "commitment", "support", "supports", "buy",
    ),
    "data": (
        "data", "dataset", "datasets", "information", "record", "records",
        "accuracy", "accurate", "accessible", "organised", "organized",
        "clean", "quality", "complete", "reliable", "structured",
    ),
    "govern": (
        "governance", "ethic", "ethics", "ethical", "privacy", "security",
        "secure", "compliance", "comply", "regulation", "regulatory",
        "legal", "law", "accountability", "accountable", "oversight",
        "transparency", "transparent", "bias", "biased", "fair", "fairness",
        "trust", "trusted", "trustworthy", "risk", "protection", "gdpr",
        "audit", "monitor", "monitoring",
    ),
    "culture": (
        "culture", "cultural", "innovation", "innovative", "experiment",
        "experimentation", "change", "resistant", "resistance", "open",
        "openness", "adapt", "adaptable", "mindset", "willing", "attitude",
        "failure", "failed", "iterate", "iteration", "continuous",
    ),
    "adopt": (
        "adopt", "adoption", "adopting", "implement", "implementation",
        "deploy", "deployment", "integrate", "rollout", "pilot", "use",
        "using", "usage", "apply", "application",
    ),
    "ai": (
        "ai", "artificial", "intelligence", "automation", "automated",
        "machine", "algorithm", "algorithmic", "model", "models", "tool",
        "tools", "technology", "technologies", "digital",
    ),
}

# Flattened for lookup: word -> the concepts it belongs to.
_WORD_TO_CONCEPTS: Dict[str, List[str]] = {}
for _concept, _words in CONCEPT_LEXICON.items():
    for _word in _words:
        _WORD_TO_CONCEPTS.setdefault(_word, []).append(_concept)


def _stem(word: str) -> str:
    """
    Crude suffix stripping.

    A full stemmer is more accurate and is another dependency. For matching
    survey wording, folding plurals and the common verb endings together is
    most of the benefit: "trains", "training" and "trained" all need to look
    like one word, and this does that.
    """
    for suffix in ("ational", "iveness", "fulness", "ousness", "ingly",
                   "edly", "ation", "ments", "ement", "ness", "ing", "ers",
                   "ies", "ied", "ive", "ity", "al", "ed", "es", "ly", "s"):
        if len(word) > len(suffix) + 3 and word.endswith(suffix):
            return word[: -len(suffix)]
    return word


def _terms(text: str) -> List[str]:
    """
    Turn a question into the terms it will be compared on.

    Each meaningful word contributes its stem, plus the concept it belongs to
    if it is a domain word. The concepts are what let differently-worded
    questions meet: "broadband" and "servers" and "infrastructure" all add
    `infra`, so a question about any of them lines up with a question about
    the others.
    """
    terms: List[str] = []
    for raw in _normalise(text).split():
        if raw in STOPWORDS or len(raw) < 2:
            continue
        stem = _stem(raw)
        terms.append(stem)
        for concept in _WORD_TO_CONCEPTS.get(raw, ()):
            terms.append(f"~{concept}")
        if raw != stem:
            for concept in _WORD_TO_CONCEPTS.get(stem, ()):
                terms.append(f"~{concept}")
    return terms


def _build_idf() -> Dict[str, float]:
    """
    Inverse document frequency across the 32 questions.

    Without this, every question scores highly against every other one purely
    because they all say "AI" and "our company". Weighting by rarity means the
    words that actually distinguish one question from another are the ones
    that decide the match.
    """
    documents = [set(_terms(item.text)) for item in cfg.ALL_ITEMS]
    total = len(documents)
    counts: Dict[str, int] = {}
    for document in documents:
        for term in document:
            counts[term] = counts.get(term, 0) + 1
    # Smoothed, so a term appearing in every question still carries a little.
    return {
        term: math.log((total + 1) / (count + 1)) + 1.0
        for term, count in counts.items()
    }


_IDF: Dict[str, float] = {}


def _idf() -> Dict[str, float]:
    global _IDF
    if not _IDF:
        _IDF = _build_idf()
    return _IDF


def _vector(text: str) -> Dict[str, float]:
    """IDF-weighted term frequencies for one piece of text."""
    idf = _idf()
    counts: Dict[str, float] = {}
    for term in _terms(text):
        counts[term] = counts.get(term, 0.0) + 1.0
    # An unseen term still gets a weight, or a question using vocabulary the
    # instrument never uses would contribute nothing at all.
    default = math.log(len(cfg.ALL_ITEMS) + 1) + 1.0
    return {term: count * idf.get(term, default) for term, count in counts.items()}


def _cosine(left: Dict[str, float], right: Dict[str, float]) -> float:
    if not left or not right:
        return 0.0
    shared = set(left) & set(right)
    if not shared:
        return 0.0
    dot = sum(left[term] * right[term] for term in shared)
    size_left = math.sqrt(sum(value * value for value in left.values()))
    size_right = math.sqrt(sum(value * value for value in right.values()))
    return dot / (size_left * size_right) if size_left and size_right else 0.0


def meaning_similarity(left: str, right: str) -> float:
    """How close two questions are in meaning, 0 to 1."""
    return _cosine(_vector(left), _vector(right))


def find_item_code(header: str) -> Optional[str]:
    """
    Look for one of the tool's own item codes anywhere in a header.

    The readable CSV this tool writes puts the code in the middle of the
    header, between the factor name and the question. Spotting it is both
    faster and safer than matching the wording around it -- and the tool not
    being able to read a file it produced itself is a poor look.
    """
    for token in re.split(r"[^A-Za-z0-9_]+", str(header).upper()):
        if token in cfg.ITEMS_BY_ID:
            return token
    return None


def _question_text(header: str) -> str:
    """
    Pull the question out of a column header.

    Google Forms writes "Section name [the actual question]". The readable CSV
    writes "Factor | CODE | question  [REVERSE-WORDED]". Other exports use
    "Section name - question", or just the question. All of them end up as the
    question alone.
    """
    header = str(header).strip()

    # Annotations in shouting capitals are labels about the question, not part
    # of it. Stripped before anything else so they can't drag a match down.
    header = re.sub(r"\[\s*[A-Z][A-Z\s_\-]{3,}\s*\]", " ", header).strip()

    # Pipe-separated headers: the question is the last and longest part.
    if "|" in header:
        parts = [part.strip() for part in header.split("|") if part.strip()]
        if parts:
            return max(parts, key=len)

    bracketed = re.search(r"\[(.+)\]", header, flags=re.DOTALL)
    if bracketed:
        return bracketed.group(1).strip()

    for separator in (" - ", " – ", " — ", ": "):
        if separator in header:
            head, _, tail = header.partition(separator)
            # Only treat it as a prefix if the left side looks like a section
            # name rather than the start of the question itself.
            if len(head) < 60 and not head.endswith("?"):
                return tail.strip()

    return header


def _similarity(left: str, right: str) -> float:
    return SequenceMatcher(None, left, right).ratio()


# ---------------------------------------------------------------------------
# What came back from an import
# ---------------------------------------------------------------------------

@dataclass
class ColumnMatch:
    header: str
    target: str
    kind: str          # "item" or "context"
    confidence: float
    how: str           # short description of why it matched


@dataclass
class ImportReport:
    """Everything the import produced, including what it couldn't place."""

    rows: List[Dict[str, object]] = field(default_factory=list)
    matches: List[ColumnMatch] = field(default_factory=list)
    unmatched_headers: List[str] = field(default_factory=list)
    free_text_headers: List[str] = field(default_factory=list)
    ignored_headers: List[str] = field(default_factory=list)
    error: Optional[str] = None

    @property
    def matched_items(self) -> List[str]:
        return [m.target for m in self.matches if m.kind == "item"]

    @property
    def matched_context(self) -> List[str]:
        return [m.target for m in self.matches if m.kind == "context"]

    @property
    def missing_items(self) -> List[str]:
        found = set(self.matched_items)
        return [item.id for item in cfg.ALL_ITEMS if item.id not in found]

    @property
    def missing_context(self) -> List[str]:
        found = set(self.matched_context)
        return [f for f in cfg.CATEGORICAL_FIELDS if f not in found]

    @property
    def item_coverage(self) -> float:
        """Share of the tool's questions this file contains, however matched."""
        return len(set(self.matched_items)) / len(cfg.ALL_ITEMS)

    @property
    def direct_matches(self) -> List[ColumnMatch]:
        """Questions recognised by their wording -- the same instrument."""
        return [m for m in self.matches if m.kind == "item" and m.how != "meaning"]

    @property
    def meaning_matches(self) -> List[ColumnMatch]:
        """
        Questions paired up by meaning rather than wording.

        These are suggestions, not conclusions. Saying that another survey's
        question measures the same thing as one of ours is a judgement about
        the constructs involved, which belongs to the researcher rather than
        to a similarity score.
        """
        return [m for m in self.matches if m.kind == "item" and m.how == "meaning"]

    @property
    def direct_coverage(self) -> float:
        """Share of questions matched on wording alone."""
        return len({m.target for m in self.direct_matches}) / len(cfg.ALL_ITEMS)

    @property
    def looks_like_a_different_survey(self) -> bool:
        """
        True when the file is a survey, but not this one.

        Judged on wording matches only. A file that lines up question for
        question is this questionnaire in another format, and can simply be
        read. A file that only lines up by meaning is a different instrument,
        and the difference matters: its scores are an approximation of this
        one's, not the same measurement.
        """
        return self.direct_coverage < 0.25

    def compare_to_framework(self) -> Dict[str, object]:
        """
        Read a different survey against the seven factors.

        When a file turns out to be another instrument, refusing it and
        stopping there wastes what is actually a useful question: how much of
        this framework does that survey cover, and what does it ask about that
        this one doesn't? That comparison is worth having for the write-up,
        and it is honest in a way that forcing a score out of the file would
        not be.

        Returns the best-matching factor for each of the file's questions,
        which factors go untouched, and which questions correspond to nothing
        here at all.
        """
        questions = [
            header for header in
            (self.unmatched_headers + self.free_text_headers
             + [m.header for m in self.meaning_matches])
            if len(str(header).split()) > 3
        ]

        per_factor: Dict[str, List[Tuple[str, float]]] = {f.id: [] for f in cfg.FACTORS}
        unrelated: List[Tuple[str, float]] = []

        for header in questions:
            text = _question_text(header)
            best_factor, best_score = "", 0.0
            for factor in cfg.FACTORS:
                score = max(
                    meaning_similarity(text, item.text) for item in factor.items
                )
                if score > best_score:
                    best_factor, best_score = factor.id, score

            # Below this, the question is not a weak match for a factor -- it
            # is about something the framework does not cover. In the survey
            # this was built against, the questions scoring here asked whether
            # AI had already improved services, which is an outcome rather
            # than a readiness condition.
            if best_score < 0.25:
                unrelated.append((header, best_score))
            else:
                per_factor[best_factor].append((header, best_score))

        return {
            "by_factor": per_factor,
            "covered": [f.id for f in cfg.FACTORS if per_factor[f.id]],
            "not_covered": [f.id for f in cfg.FACTORS if not per_factor[f.id]],
            "unrelated": unrelated,
            "questions_examined": len(questions),
        }


# ---------------------------------------------------------------------------
# Reading the file
# ---------------------------------------------------------------------------

def read_table(data: bytes, filename: str = "") -> Tuple[List[str], List[Dict[str, object]]]:
    """
    Read a CSV or Excel file into headers and rows.

    Excel is supported because that is what people have. Asking someone to
    convert to CSV first is a step that only exists to suit the code.
    """
    name = (filename or "").lower()

    if name.endswith((".xlsx", ".xlsm", ".xltx")):
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        table = [list(row) for row in sheet.iter_rows(values_only=True)]
        workbook.close()

        if not table:
            return [], []
        headers = [str(h).strip() if h is not None else "" for h in table[0]]
        rows = [
            {headers[i]: row[i] for i in range(min(len(headers), len(row)))}
            for row in table[1:]
            if any(cell is not None and str(cell).strip() for cell in row)
        ]
        return headers, rows

    # Default to CSV. utf-8-sig drops the byte-order mark Excel likes to add.
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - latin-1 decodes anything
        raise ValueError("Could not decode the file as text.")

    reader = csv.DictReader(io.StringIO(text))
    headers = [h.strip() for h in (reader.fieldnames or [])]
    rows = [dict(row) for row in reader]
    return headers, rows


# ---------------------------------------------------------------------------
# Working out which column is which
# ---------------------------------------------------------------------------

def match_columns(headers: Sequence[str]) -> Tuple[List[ColumnMatch], List[str], List[str], List[str]]:
    """
    Decide what each column is.

    Returns the matches, plus the headers that were left over, split into
    likely free-text questions, ignored admin columns, and the rest.
    """
    matches: List[ColumnMatch] = []
    leftover: List[str] = []
    free_text: List[str] = []
    ignored: List[str] = []

    item_texts = {item.id: _normalise(item.text) for item in cfg.ALL_ITEMS}

    # Every plausible (header, target) pairing with its score, not just the
    # single best per header. Two questions can word similarly enough that
    # they both pick the same item first; keeping the runners-up lets the
    # loser fall back to its own question instead of being thrown away.
    proposals: List[Tuple[str, str, str, float, str]] = []
    considered: set = set()

    for header in headers:
        if not str(header).strip():
            continue

        flat = _normalise(header)

        if flat in IGNORED_HEADERS:
            ignored.append(header)
            continue

        considered.add(header)

        # 1. The column is, or contains, one of the tool's item codes.
        code = find_item_code(header)
        if code:
            proposals.append((header, code, "item", 1.0, "item code"))
            continue

        # 2. Both readings are proposed and the strongest wins later, rather
        # than the first check that fires taking the column. A question can
        # easily contain a context word -- "We provide training opportunities
        # to employees..." is not the employee-count column -- so letting a
        # context check short-circuit here silently loses real questions.
        found_any = False

        raw_question = _question_text(header)
        question = _normalise(raw_question)
        for item_id, text in item_texts.items():
            score = _similarity(question, text)
            if score >= TEXT_MATCH_THRESHOLD:
                how = "question text" if score < 0.999 else "question text (exact)"
                proposals.append((header, item_id, "item", score, how))
                found_any = True

        context_hit, context_score = _best_context(flat)
        if context_hit and context_score >= 0.80:
            how = ("context column" if context_score >= 0.90
                   else "context column (loose match)")
            proposals.append((header, context_hit, "context", context_score, how))
            found_any = True

        if found_any:
            continue

        # 3. Nothing matched on wording. Try meaning, for a question asking
        # the same thing in someone else's words. These score lower than a
        # wording match by nature, so they are proposed below every lexical
        # one and marked so they can be confirmed rather than trusted.
        by_meaning = sorted(
            ((item_id, meaning_similarity(raw_question, cfg.ITEMS_BY_ID[item_id].text))
             for item_id in item_texts),
            key=lambda pair: -pair[1],
        )
        for item_id, score in by_meaning[:3]:
            if score >= MEANING_MATCH_THRESHOLD:
                # Scaled below the wording threshold so a lexical match on any
                # other column always outranks a meaning match on this one.
                proposals.append((header, item_id, "item",
                                  TEXT_MATCH_THRESHOLD * score, "meaning"))
                found_any = True

        if found_any:
            continue

        # 3. Not part of the assessment. An open question is worth keeping.
        if "?" in str(header) or len(str(header).split()) > 6:
            free_text.append(header)
        else:
            leftover.append(header)

    # Greedy assignment, strongest pairing first. A header and a target are
    # each used once, so the best overall reading of the file wins rather than
    # whichever column happened to come first.
    used_headers: set = set()
    used_targets: set = set()
    for header, target, kind, score, how in sorted(proposals, key=lambda p: -p[3]):
        if header in used_headers or target in used_targets:
            continue
        used_headers.add(header)
        used_targets.add(target)
        matches.append(ColumnMatch(header, target, kind, score, how))

    placed = used_headers | set(free_text) | set(leftover)
    leftover.extend(h for h in considered if h not in placed)

    return matches, leftover, free_text, ignored


def _best_context(flat_header: str) -> Tuple[str, float]:
    best_field, best_score = "", 0.0
    for field_name, aliases in CONTEXT_ALIASES.items():
        for alias in aliases:
            alias_flat = _normalise(alias)
            if alias_flat == flat_header:
                return field_name, 1.0
            score = _similarity(flat_header, alias_flat)
            # A header that is basically the alias plus a little decoration is
            # a strong signal, e.g. "Q2. Country selection (required)". The
            # length guard matters: without it a single word like "employees"
            # appearing anywhere in a long question would capture the column.
            if (
                alias_flat in flat_header
                and len(alias_flat) >= 6
                and len(flat_header) <= len(alias_flat) * 2.5
            ):
                score = max(score, 0.93)
            if score > best_score:
                best_field, best_score = field_name, score
    return best_field, best_score


# ---------------------------------------------------------------------------
# Putting it together
# ---------------------------------------------------------------------------

def import_survey(data: bytes, filename: str = "",
                  use_meaning_matches: bool = False) -> ImportReport:
    """
    Read a survey export and map it onto the tool's schema.

    Always returns a report. A file that couldn't be read comes back with
    `error` set rather than raising, because the caller is a web page and a
    stack trace is not an answer.
    """
    try:
        headers, raw_rows = read_table(data, filename)
    except Exception as problem:  # noqa: BLE001 - shown to the user verbatim
        return ImportReport(error=str(problem))

    if not headers:
        return ImportReport(error="The file has no header row.")
    if not raw_rows:
        return ImportReport(error="The file has a header row but no responses.")

    matches, leftover, free_text, ignored = match_columns(headers)

    # A meaning match is a suggestion. On a file that is plainly this
    # questionnaire, one reworded question among thirty-one recognised ones is
    # safe to accept. On a file that is a different instrument, the same
    # suggestion would be quietly inventing data, so the values are left out
    # and the matches are kept for analysis only.
    direct = {m.target for m in matches if m.kind == "item" and m.how != "meaning"}
    foreign = len(direct) / len(cfg.ALL_ITEMS) < 0.25
    usable = [
        m for m in matches
        if not (foreign and m.kind == "item" and m.how == "meaning"
                and not use_meaning_matches)
    ]

    rows: List[Dict[str, object]] = []
    for index, raw in enumerate(raw_rows, start=1):
        row: Dict[str, object] = {"company_id": f"R{index:04d}"}

        for match in usable:
            value = raw.get(match.header)
            if match.kind == "item":
                parsed = cfg.parse_rating(value)
                if parsed is not None:
                    row[match.target] = parsed
            else:
                tidied = cfg.normalise_category(match.target, value)
                if tidied is not None:
                    row[match.target] = tidied

        # Keep the open-ended answers. They are never scored, but they are
        # where respondents explain the numbers, so throwing them away loses
        # the most useful part of the response.
        notes: Dict[str, str] = {}
        for header in free_text:
            text = raw.get(header)
            if text is not None and str(text).strip():
                notes[str(header).strip()] = str(text).strip()
        if notes:
            row["free_text"] = notes

        rows.append(row)

    return ImportReport(
        rows=rows,
        matches=matches,
        unmatched_headers=leftover,
        free_text_headers=free_text,
        ignored_headers=ignored,
    )


def profiles_from_report(report: ImportReport):
    """Turn an import report into (identifier, CompanyProfile) pairs."""
    from .company_profile import CompanyProfile

    pairs = []
    for row in report.rows:
        identifier = str(row.get("company_id") or f"row{len(pairs) + 1}")
        pairs.append((identifier, CompanyProfile.from_dict(row)))
    return pairs
