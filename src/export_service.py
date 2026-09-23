"""
Export service

Turns a finished AssessmentResults into something the company can take away:
a PDF report to read, or an Excel workbook to dig into.

Everything here works in memory and returns bytes. That's deliberate -- the
dashboard hands those bytes straight to a download button, and nothing ever
touches the server's disk, which is what keeps the no-database, nothing-stored
design intact. The write_pdf/write_excel helpers exist for testing and for
running from the command line.

Note on what goes in the file: the export carries the company's answers and
scores but no company name, contact details or anything else identifying,
because CompanyProfile never collects those in the first place.
"""

from __future__ import annotations

import io
import math
from datetime import datetime
from typing import Any, Dict, List, Optional, Sequence, Tuple
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.chart import BarChart, RadarChart, Reference
from openpyxl.chart.data_source import AxDataSource, StrRef
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.graphics.shapes import Circle, Drawing, Line, Polygon, String
from reportlab.platypus import (
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from . import score_configuration as cfg
from .assessment_results import AssessmentResults, FactorScore, ItemScore


# ---------------------------------------------------------------------------
# Shared look-up tables
# ---------------------------------------------------------------------------
# One place for the colours so the PDF and the spreadsheet agree with each
# other. If a band colour changes it changes in both.

BAND_COLOURS: Dict[str, str] = {
    "Emerging": "C0392B",    # red
    "Developing": "D68910",  # amber
    "Advanced": "1E8449",    # green
}

SEVERITY_COLOURS: Dict[str, str] = {
    "critical": "C0392B",
    "moderate": "D68910",
    "refine": "2874A6",
    # Raised by the written answers rather than by a score crossing a
    # threshold. Purple keeps it visually distinct from the scored severities,
    # because the evidence behind it is a different kind of evidence.
    "raised": "7D3C98",
}

SEVERITY_LABELS: Dict[str, str] = {
    "critical": "Critical",
    "moderate": "Moderate",
    "refine": "Refine",
    "raised": "Raised in comments",
}

# Nicer headings than the raw field names for the context block.
# Full factor names don't fit round a radar, so each gets a one-word label.
SHORT_FACTOR_LABELS: Dict[str, str] = {
    "budget": "Budget", "workforce": "Workforce", "leadership": "Leadership",
    "data": "Data", "technology": "Technology", "culture": "Culture",
    "governance": "Governance",
}

CONTEXT_LABELS: Dict[str, str] = {
    "industry_sector": "Industry sector",
    "employee_band": "Company size",
    "years_in_operation": "Years in operation",
    "region": "Region",
    "current_ai_stage": "Current AI adoption stage",
}

# Excel mangles "1-9" and "10-49" into dates when the sheet is opened, the same
# way it does with a CSV. The workbook writes these spelled-out forms instead.
EXCEL_SAFE_SIZES: Dict[str, str] = {
    "1-9": "1 to 9 employees",
    "10-49": "10 to 49 employees",
    "50-249": "50 to 249 employees",
    "250+": "250+ employees",
}


def _safe_size(value: Optional[str]) -> Optional[str]:
    return EXCEL_SAFE_SIZES.get(str(value), value)


def _timestamp() -> str:
    return datetime.now().strftime("%d %B %Y, %H:%M")


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def _styles() -> Dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "title", parent=base["Title"], fontSize=20, spaceAfter=2, leading=24,
        ),
        "subtitle": ParagraphStyle(
            "subtitle", parent=base["Normal"], fontSize=9.5,
            textColor=colors.HexColor("#666666"), spaceAfter=14,
        ),
        "h2": ParagraphStyle(
            "h2", parent=base["Heading2"], fontSize=13, spaceBefore=16,
            spaceAfter=7, textColor=colors.HexColor("#1A1A1A"),
        ),
        "body": ParagraphStyle(
            "body", parent=base["Normal"], fontSize=9.5, leading=13.5,
            alignment=TA_LEFT,
        ),
        "small": ParagraphStyle(
            "small", parent=base["Normal"], fontSize=8.5, leading=11.5,
            textColor=colors.HexColor("#555555"),
        ),
        "rec_title": ParagraphStyle(
            "rec_title", parent=base["Normal"], fontSize=10, leading=13,
            fontName="Helvetica-Bold",
        ),
    }


def _score_bar(score: float, width: float, colour: str) -> Table:
    """
    A horizontal bar drawn as a two-cell table.

    Using a table rather than reportlab's chart classes keeps it simple and it
    flows with the rest of the document instead of needing absolute placement.
    """
    filled = max(0.0, min(1.0, score / 100.0)) * width
    empty = width - filled

    # A zero-width column still draws a hairline, so drop the empty cell when
    # the score is full, and the filled cell when it's zero.
    if filled <= 0.01:
        widths, cells = [width], [[""]]
        style = [("BACKGROUND", (0, 0), (0, 0), colors.HexColor("#E8E8E8"))]
    elif empty <= 0.01:
        widths, cells = [width], [[""]]
        style = [("BACKGROUND", (0, 0), (0, 0), colors.HexColor("#" + colour))]
    else:
        widths, cells = [filled, empty], [["", ""]]
        style = [
            ("BACKGROUND", (0, 0), (0, 0), colors.HexColor("#" + colour)),
            ("BACKGROUND", (1, 0), (1, 0), colors.HexColor("#E8E8E8")),
        ]

    bar = Table(cells, colWidths=widths, rowHeights=[5 * mm])
    bar.setStyle(TableStyle(style + [
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
    ]))
    return bar


def _evidence_markup(rec, single: bool = False) -> str:
    """
    The respondents' own words, rendered under the recommendation they support.

    An action carries more weight when the reader can see it agrees with what
    someone actually wrote, and it is worth seeing when it does not. The quotes
    are evidence for the recommendation, never an input to the score.
    """
    if not getattr(rec, "evidence", None):
        return ""

    if single:
        lead = "In your own words:"
    else:
        people = "respondent" if rec.evidence_count == 1 else "respondents"
        lead = f"Raised in writing by {rec.evidence_count} {people}:"

    quotes = "".join(
        f'<br/><i>&ldquo;{escape(quote)}&rdquo;</i>' for quote in rec.evidence
    )
    return (f'<br/><br/><font size="8" color="#777777">{lead}{quotes}</font>')


def _radar_drawing(scores: List[Tuple[str, float]], colour: str,
                   size: float = 175.0) -> Drawing:
    """
    The profile shape, drawn for the PDF.

    The bar chart ranks the factors; this shows whether readiness is even or
    lopsided, which is a different question and the one people tend to ask
    first when they see the report. The dashboard shows it, so the report
    people take away should show it too.
    """
    drawing = Drawing(size, size)
    centre = size / 2.0
    radius = size * 0.32
    count = len(scores)

    def point(index: int, value: float) -> Tuple[float, float]:
        # Clockwise from the top. reportlab's y axis points up where the
        # dashboard's SVG points down, so the sign is flipped here to keep
        # the report and the screen showing the same shape the same way round.
        angle = (math.pi / 2) - (2 * math.pi * index / count)
        distance = radius * max(0.0, min(100.0, value)) / 100.0
        return (centre + distance * math.cos(angle),
                centre + distance * math.sin(angle))

    for fraction in (0.25, 0.5, 0.75, 1.0):
        drawing.add(Circle(centre, centre, radius * fraction,
                           fillColor=None,
                           strokeColor=colors.HexColor("#E4E4E4"),
                           strokeWidth=0.5))

    for index, (name, _) in enumerate(scores):
        end_x, end_y = point(index, 100)
        drawing.add(Line(centre, centre, end_x, end_y,
                         strokeColor=colors.HexColor("#E4E4E4"), strokeWidth=0.5))

        label_x, label_y = point(index, 133)
        label = String(label_x, label_y - 2, name, fontSize=6.2,
                       fillColor=colors.HexColor("#666666"))
        if label_x < centre - 6:
            label.textAnchor = "end"
        elif label_x > centre + 6:
            label.textAnchor = "start"
        else:
            label.textAnchor = "middle"
        drawing.add(label)

    outline: List[float] = []
    for index, (_, value) in enumerate(scores):
        x, y = point(index, value)
        outline.extend([x, y])

    drawing.add(Polygon(outline,
                        fillColor=colors.HexColor("#" + colour),
                        fillOpacity=0.22,
                        strokeColor=colors.HexColor("#" + colour),
                        strokeWidth=1.4))
    return drawing


def to_pdf_bytes(results: AssessmentResults) -> bytes:
    """Build the PDF report and return it as bytes."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title="AI Adoption Readiness Assessment",
        author="AI Adoption Readiness Assessment Tool",
    )

    s = _styles()
    content_width = doc.width
    story: List = []

    # --- header -----------------------------------------------------------
    story.append(Paragraph("AI Adoption Readiness Assessment", s["title"]))
    story.append(Paragraph(f"Generated {_timestamp()}", s["subtitle"]))

    # --- headline score ---------------------------------------------------
    band_colour = BAND_COLOURS.get(results.readiness_tier, "555555")
    headline = Table(
        [[
            Paragraph(
                f'<font size="30"><b>{results.overall_score:.1f}</b></font>'
                f'<font size="12" color="#777777">/100</font>',
                s["body"],
            ),
            Paragraph(
                f'<font size="15" color="#{band_colour}"><b>'
                f'{escape(results.readiness_tier)}</b></font><br/><br/>'
                f'{escape(results.tier_description)}',
                s["body"],
            ),
        ]],
        colWidths=[38 * mm, content_width - 38 * mm],
    )
    headline.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (0, 0), 0),
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F7F7F7")),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ("LEFTPADDING", (0, 0), (-1, -1), 10),
        ("RIGHTPADDING", (0, 0), (-1, -1), 10),
        ("LINEBEFORE", (0, 0), (0, 0), 3, colors.HexColor("#" + band_colour)),
    ]))
    story.append(headline)

    # --- company context --------------------------------------------------
    context_pairs: List[str] = []
    for field, label in CONTEXT_LABELS.items():
        value = results.context.get(field)
        if value:
            shown = _safe_size(value) if field == "employee_band" else value
            context_pairs.append(f"<b>{escape(label)}:</b> {escape(str(shown))}")
    if context_pairs:
        story.append(Spacer(1, 7))
        story.append(Paragraph("&nbsp;&nbsp;|&nbsp;&nbsp;".join(context_pairs), s["small"]))

    # --- how to read it ---------------------------------------------------
    # Placed before the scores rather than after them: the reader meets a
    # number out of 100 on the next line, and needs to know what produced it
    # at that moment, not on the last page.
    story.append(Paragraph("How the score is calculated", s["h2"]))
    story.append(Paragraph(
        f"Each of the {len(cfg.FACTORS)} factors is scored from the answers "
        f"given to its questions, on a 1&ndash;5 agreement scale, converted to "
        f"a 0&ndash;100 scale. Questions worded as barriers are reversed first, "
        f"so a higher score always means more ready. The overall score is the "
        f"weighted average of the factor scores. Bands: Emerging 0&ndash;40, "
        f"Developing 40&ndash;70, Advanced 70&ndash;100.",
        s["small"],
    ))

    # --- factor scores ----------------------------------------------------
    story.append(Paragraph("Readiness by factor", s["h2"]))

    label_w = 62 * mm
    score_w = 16 * mm
    band_w = 22 * mm
    bar_w = content_width - label_w - score_w - band_w

    rows: List[List] = []
    for factor in results.ordered_factors():
        colour = BAND_COLOURS.get(factor.band_label, "555555")
        rows.append([
            Paragraph(escape(factor.name), s["body"]),
            Paragraph(f"<b>{factor.score:.1f}</b>", s["body"]),
            _score_bar(factor.score, bar_w - 4 * mm, colour),
            Paragraph(
                f'<font color="#{colour}">{escape(factor.band_label)}</font>',
                s["small"],
            ),
        ])

    factor_table = Table(rows, colWidths=[label_w, score_w, bar_w, band_w])
    factor_table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (0, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LINEBELOW", (0, 0), (-1, -2), 0.4, colors.HexColor("#EAEAEA")),
    ]))
    story.append(factor_table)

    # --- profile shape ----------------------------------------------------
    ordered = results.ordered_factors()
    shape = _radar_drawing(
        [(SHORT_FACTOR_LABELS.get(f.factor_id, f.name), f.score) for f in ordered],
        band_colour,
    )
    shape_row = Table(
        [[shape, Paragraph(
            "<b>Profile shape</b><br/><br/>"
            "An even shape means readiness is spread across the seven areas. "
            "A lopsided one means some areas are far ahead of others, which "
            "usually matters more than the overall number: adoption tends to "
            "be held back by the weakest area rather than helped by the "
            "strongest.",
            s["small"])]],
        colWidths=[76 * mm, content_width - 76 * mm],
    )
    shape_row.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (0, 0), 0),
    ]))
    story.append(Spacer(1, 6))
    story.append(shape_row)

    # --- strengths and barriers ------------------------------------------
    if results.strengths:
        story.append(Paragraph("Strengths", s["h2"]))
        for factor in results.strengths:
            story.append(Paragraph(
                f"<b>{escape(factor.name)}</b> &mdash; {factor.score:.0f}/100",
                s["body"],
            ))

    if results.barriers:
        story.append(Paragraph("Main barriers", s["h2"]))
        for factor in results.barriers:
            story.append(Paragraph(
                f"<b>{escape(factor.name)}</b> &mdash; {factor.score:.0f}/100",
                s["body"],
            ))

    if results.item_barriers:
        story.append(Paragraph("Specific gaps", s["h2"]))
        gap_rows = []
        for item in results.item_barriers:
            # Reverse-worded questions describe a problem, so agreeing with one
            # is what produces the low score. Without saying so, the statement
            # and the score next to it look like they contradict each other.
            note = (
                ' <i><font color="#888888">(barrier statement &mdash; agreeing '
                'lowers readiness)</font></i>'
                if item.reverse else ""
            )
            gap_rows.append([
                Paragraph(f"<b>{escape(item.item_id)}</b>", s["small"]),
                Paragraph(escape(item.text) + note, s["small"]),
                Paragraph(f"{item.score:.0f}/100", s["small"]),
            ])

        gap_table = Table(gap_rows, colWidths=[18 * mm, content_width - 40 * mm, 22 * mm])
        gap_table.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (0, -1), 0),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        story.append(gap_table)

    # --- recommendations --------------------------------------------------
    story.append(Paragraph("Recommended actions", s["h2"]))

    if not results.recommendations:
        story.append(Paragraph(
            "No actions are flagged. Every factor scored at or above the "
            "Advanced threshold.",
            s["body"],
        ))
    else:
        story.append(Paragraph(
            "Listed in priority order. Priority reflects how much closing each "
            "gap would move the overall readiness score.",
            s["small"],
        ))
        story.append(Spacer(1, 8))

        for index, rec in enumerate(results.recommendations, start=1):
            colour = SEVERITY_COLOURS.get(rec.severity, "555555")
            label = SEVERITY_LABELS.get(rec.severity, rec.severity.title())
            source = rec.factor_name
            if rec.triggered_by_item:
                source += f" &middot; {escape(rec.triggered_by_item)}"

            block = Table(
                [[Paragraph(
                    f'<font color="#{colour}"><b>{index}. {escape(label)}</b></font>'
                    f'&nbsp;&nbsp;<font color="#777777" size="8">{source}</font><br/>'
                    f'<b>{escape(rec.title)}</b><br/>'
                    f'{escape(rec.action)}'
                    f'{_evidence_markup(rec, single=True)}',
                    s["body"],
                )]],
                colWidths=[content_width],
            )
            block.setStyle(TableStyle([
                ("LINEBEFORE", (0, 0), (0, 0), 2.5, colors.HexColor("#" + colour)),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FAFAFA")),
            ]))
            story.append(KeepTogether(block))
            story.append(Spacer(1, 6))

    # --- every question ---------------------------------------------------
    # The factor scores say which area is weak. This says which question
    # inside it is, which is the level an action can actually be taken at.
    story.append(Paragraph("Every question", s["h2"]))
    story.append(Paragraph(
        "All answers, grouped by factor. Barrier-worded questions are marked "
        "&mdash; on those, agreeing lowers readiness, and the score shown is "
        "after that has been accounted for.",
        s["small"],
    ))
    story.append(Spacer(1, 5))

    detail_rows: List[List] = []
    for factor in ordered:
        detail_rows.append([
            Paragraph(f"<b>{escape(factor.name)}</b>", s["small"]),
            Paragraph(f"<b>{factor.score:.0f}</b>", s["small"]),
            "",
        ])
        for item in factor.item_scores:
            marker = " <font color='#999999'>(barrier-worded)</font>" if item.reverse else ""
            band = "Advanced" if item.score >= 70 else (
                "Developing" if item.score >= 40 else "Emerging")
            detail_rows.append([
                Paragraph(
                    f"<font color='#888888'>{escape(item.item_id)}</font> "
                    f"{escape(item.text)}{marker}", s["small"]),
                Paragraph(f"{item.score:.0f}", s["small"]),
                _score_bar(item.score, 28 * mm, BAND_COLOURS.get(band, "555555")),
            ])

    detail = Table(detail_rows,
                   colWidths=[content_width - 46 * mm, 12 * mm, 34 * mm])
    detail.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (0, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 2.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
    ]))
    story.append(detail)

    # --- what they wrote --------------------------------------------------
    if results.notes:
        story.append(Paragraph("In your own words", s["h2"]))
        story.append(Paragraph(
            "Not scored, and deliberately so &mdash; but this is usually where "
            "the reason behind a low score sits.",
            s["small"],
        ))
        story.append(Spacer(1, 4))
        for question, answer in results.notes.items():
            story.append(Paragraph(f"<b>{escape(question)}</b>", s["small"]))
            story.append(Paragraph(escape(answer), s["body"]))
            story.append(Spacer(1, 5))

        insights = results.text_insights
        if insights is not None and insights.by_factor:
            named = ", ".join(escape(sig.factor_name)
                              for sig in insights.by_factor)
            story.append(Paragraph(
                f"<b>Read as being about:</b> {named}. Where one of these was "
                f"not already flagged by the scores, it appears in the actions "
                f"above marked as raised in comments.", s["small"]))
        elif insights is not None and insights.answers_dropped:
            story.append(Paragraph(
                f"{insights.answers_dropped} of {insights.answers_seen} "
                f"written answers said nothing specific and were not used.",
                s["small"]))

    story.append(Spacer(1, 10))
    story.append(Paragraph(
        "This assessment is based on self-reported answers and is intended to "
        "support planning, not to replace professional advice.",
        s["small"],
    ))

    doc.build(story)
    return buffer.getvalue()


def write_pdf(results: AssessmentResults, path: str) -> str:
    """Write the PDF report to disk. Returns the path."""
    with open(path, "wb") as handle:
        handle.write(to_pdf_bytes(results))
    return path


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

_HEADER_FILL = PatternFill("solid", fgColor="2C3E50")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_TITLE_FONT = Font(bold=True, size=14)
_THIN = Side(style="thin", color="D5D5D5")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _write_header(sheet, row: int, headers: List[str]) -> None:
    for column, heading in enumerate(headers, start=1):
        cell = sheet.cell(row=row, column=column, value=heading)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[row].height = 26


def _set_widths(sheet, widths: List[Tuple[int, int]]) -> None:
    for column, width in widths:
        sheet.column_dimensions[get_column_letter(column)].width = width


def _label_chart(chart, categories: Reference, *,
                 category_title: Optional[str] = None,
                 value_title: Optional[str] = None,
                 show_values: bool = True,
                 value_axis: bool = True,
                 value_range: Optional[Tuple[float, float]] = None) -> None:
    """
    Give a chart its category labels, axis titles and value labels.

    The category part is not cosmetic. openpyxl's set_categories() always
    writes the range as a *numeric* reference, and the categories here are
    words -- factor names, item codes, country names. Excel reads a numeric
    reference to a column of text as empty, which is why a chart built the
    plain way comes out with no labels round the edge at all. Rewriting the
    reference as a string reference is what puts the names back.

    The axes are also marked explicitly as not deleted, because an axis left
    unspecified is not guaranteed to be drawn.
    """
    chart.set_categories(categories)

    source = AxDataSource(strRef=StrRef(f=categories))
    for series in chart.series:
        series.cat = source

    if category_title:
        chart.x_axis.title = category_title
    if value_title:
        chart.y_axis.title = value_title

    chart.x_axis.delete = False
    # A radar draws its value axis as a column of numbers straight down the
    # middle of the shape, on top of the fill. Hiding the axis removes those
    # while leaving the rings, which are a separate element and are what
    # actually convey the scale -- the range is in the chart title instead.
    chart.y_axis.delete = not value_axis

    if value_range is not None:
        # A score chart that starts at the lowest value present makes a six
        # point difference look like the whole range. Fixing the axis to
        # 0-100 keeps the bars proportional to the scores they represent.
        chart.y_axis.scaling.min, chart.y_axis.scaling.max = value_range

    if show_values:
        # Every flag is set explicitly. Left unset, some readers take "not
        # specified" as "show", and the label becomes
        # "Budget & Financial Readiness; Mean score; 55.9" sprawled across
        # its neighbours. Only the number is wanted -- the category is
        # already on the axis.
        labels = DataLabelList()
        labels.showVal = True
        labels.showSerName = False
        labels.showCatName = False
        labels.showLegendKey = False
        labels.showPercent = False
        labels.showBubbleSize = False
        chart.dataLabels = labels


def to_excel_bytes(results: AssessmentResults) -> bytes:
    """Build the Excel workbook and return it as bytes."""
    workbook = Workbook()

    # --- Summary ----------------------------------------------------------
    summary = workbook.active
    summary.title = "Summary"
    _set_widths(summary, [(1, 32), (2, 62)])

    summary["A1"] = "AI Adoption Readiness Assessment"
    summary["A1"].font = _TITLE_FONT
    summary["A2"] = f"Generated {_timestamp()}"
    summary["A2"].font = Font(size=9, color="777777")

    row = 4
    summary.cell(row=row, column=1, value="Overall readiness score").font = Font(bold=True)
    score_cell = summary.cell(row=row, column=2, value=round(results.overall_score, 1))
    score_cell.font = Font(bold=True, size=13,
                           color=BAND_COLOURS.get(results.readiness_tier, "000000"))
    row += 1

    summary.cell(row=row, column=1, value="Readiness tier").font = Font(bold=True)
    tier_cell = summary.cell(row=row, column=2, value=results.readiness_tier)
    tier_cell.font = Font(bold=True,
                          color=BAND_COLOURS.get(results.readiness_tier, "000000"))
    row += 1

    summary.cell(row=row, column=1, value="What that means").font = Font(bold=True)
    meaning = summary.cell(row=row, column=2, value=results.tier_description)
    meaning.alignment = Alignment(wrap_text=True, vertical="top")
    summary.row_dimensions[row].height = 46
    row += 2

    summary.cell(row=row, column=1, value="Company context").font = _TITLE_FONT
    row += 1
    for field, label in CONTEXT_LABELS.items():
        value = results.context.get(field)
        if not value:
            continue
        summary.cell(row=row, column=1, value=label)
        shown = _safe_size(value) if field == "employee_band" else value
        summary.cell(row=row, column=2, value=shown)
        row += 1

    row += 1
    if results.strengths:
        summary.cell(row=row, column=1, value="Strengths").font = _TITLE_FONT
        row += 1
        for factor in results.strengths:
            summary.cell(row=row, column=1, value=factor.name)
            summary.cell(row=row, column=2, value=f"{factor.score:.0f}/100")
            row += 1
        row += 1

    if results.barriers:
        summary.cell(row=row, column=1, value="Main barriers").font = _TITLE_FONT
        row += 1
        for factor in results.barriers:
            summary.cell(row=row, column=1, value=factor.name)
            summary.cell(row=row, column=2, value=f"{factor.score:.0f}/100")
            row += 1

    # --- Factor scores ----------------------------------------------------
    factors = workbook.create_sheet("Factor scores")
    _set_widths(factors, [(1, 40), (2, 12), (3, 14), (4, 10), (5, 14), (6, 20),
                          (7, 14)])
    _write_header(factors, 1, [
        "Factor", "Score /100", "Mean (1-5)", "Weight",
        "Contribution", "Band", "Short label",
    ])

    for index, factor in enumerate(results.ordered_factors(), start=2):
        factors.cell(row=index, column=1, value=factor.name)
        factors.cell(row=index, column=2, value=round(factor.score, 1))
        factors.cell(row=index, column=3, value=round(factor.mean_rating, 2))
        factors.cell(row=index, column=4, value=round(factor.weight, 4))
        factors.cell(row=index, column=5, value=round(factor.weighted_contribution, 2))
        band = factors.cell(row=index, column=6, value=factor.band_label)
        band.font = Font(bold=True,
                         color=BAND_COLOURS.get(factor.band_label, "000000"))
        # Used as the radar's category labels: the full names overlap round a
        # radar, and the column makes the abbreviation traceable.
        factors.cell(row=index, column=7,
                     value=SHORT_FACTOR_LABELS.get(factor.factor_id,
                                                   factor.factor_id))
        for column in range(1, 8):
            factors.cell(row=index, column=column).border = _BORDER

    total_row = len(results.factor_scores) + 2
    factors.cell(row=total_row, column=1, value="Overall").font = Font(bold=True)
    factors.cell(row=total_row, column=2,
                 value=round(results.overall_score, 1)).font = Font(bold=True)
    factors.cell(row=total_row, column=6,
                 value=results.readiness_tier).font = Font(bold=True)
    factors.freeze_panes = "A2"

    # A chart, because a workbook of bare numbers is worse than the PDF at the
    # one thing a spreadsheet should be good at. Two views: the factor bars,
    # and a radar that shows the shape of the profile at a glance.
    last = len(results.factor_scores) + 1
    labels = Reference(factors, min_col=1, min_row=2, max_row=last)
    values = Reference(factors, min_col=2, min_row=1, max_row=last)
    short_labels = Reference(factors, min_col=7, min_row=2, max_row=last)

    bars = BarChart()
    bars.type = "bar"
    bars.title = "Readiness by factor (0-100)"
    bars.add_data(values, titles_from_data=True)
    _label_chart(bars, labels,
                 category_title="Readiness factor",
                 value_title="Score /100", value_range=(0, 100))
    bars.height, bars.width = 9, 20
    bars.legend = None
    factors.add_chart(bars, f"A{last + 3}")

    shape = RadarChart()
    shape.type = "filled"
    shape.title = "Profile shape (0-100 by factor)"
    shape.add_data(values, titles_from_data=True)
    # No value labels on the radar -- a filled radar draws them on top of the
    # fill, and the bar chart beside it already carries the numbers. The
    # category labels are the point here.
    _label_chart(shape, short_labels, show_values=False, value_axis=False)
    shape.height, shape.width = 11, 11
    shape.y_axis.scaling.min = 0
    shape.y_axis.scaling.max = 100
    factors.add_chart(shape, f"H{last + 3}")

    # --- Responses --------------------------------------------------------
    # Both the raw answer and the re-coded one, because the difference between
    # them on reverse-worded questions is the thing people query most.
    responses = workbook.create_sheet("Responses")
    _set_widths(responses, [(1, 10), (2, 34), (3, 62), (4, 12), (5, 16), (6, 14), (7, 12)])
    _write_header(responses, 1, [
        "Item", "Factor", "Question", "Answer (1-5)",
        "Answer label", "Reverse worded?", "Score /100",
    ])

    row = 2
    for factor in results.ordered_factors():
        for item in factor.item_scores:
            responses.cell(row=row, column=1, value=item.item_id)
            responses.cell(row=row, column=2, value=item.factor_name)
            question = responses.cell(row=row, column=3, value=item.text)
            question.alignment = Alignment(wrap_text=True, vertical="top")
            responses.cell(row=row, column=4, value=item.raw_value)
            responses.cell(row=row, column=5,
                           value=cfg.RATING_LABELS.get(item.raw_value, ""))
            responses.cell(row=row, column=6, value="Yes" if item.reverse else "No")
            responses.cell(row=row, column=7, value=round(item.score, 1))
            for column in range(1, 8):
                responses.cell(row=row, column=column).border = _BORDER
            row += 1
    responses.freeze_panes = "A2"

    # Every question on one chart. This is the view that shows a factor with a
    # decent average hiding one weak question inside it, which the factor-level
    # numbers alone can't show.
    item_rows = row - 1
    item_chart = BarChart()
    item_chart.type = "bar"
    item_chart.title = "Every question, scored 0-100"
    item_chart.add_data(
        Reference(responses, min_col=7, min_row=1, max_row=item_rows),
        titles_from_data=True,
    )
    _label_chart(
        item_chart,
        Reference(responses, min_col=1, min_row=2, max_row=item_rows),
        category_title="Question",
        value_title="Score /100",
        value_range=(0, 100),
    )
    item_chart.height, item_chart.width = 20, 20
    item_chart.legend = None
    responses.add_chart(item_chart, f"I2")

    # --- Recommendations --------------------------------------------------
    recs = workbook.create_sheet("Recommendations")
    _set_widths(recs, [(1, 9), (2, 18), (3, 32), (4, 40), (5, 70), (6, 14),
                       (7, 60)])
    _write_header(recs, 1, [
        "Priority", "Severity", "Factor", "Action", "Detail", "Triggered by",
        "Supporting words from the response",
    ])

    if results.recommendations:
        for index, rec in enumerate(results.recommendations, start=1):
            row = index + 1
            recs.cell(row=row, column=1, value=index)
            severity = recs.cell(row=row, column=2,
                                 value=SEVERITY_LABELS.get(rec.severity, rec.severity))
            severity.font = Font(bold=True,
                                 color=SEVERITY_COLOURS.get(rec.severity, "000000"))
            recs.cell(row=row, column=3, value=rec.factor_name)
            recs.cell(row=row, column=4, value=rec.title)
            detail = recs.cell(row=row, column=5, value=rec.action)
            detail.alignment = Alignment(wrap_text=True, vertical="top")
            recs.cell(row=row, column=6, value=rec.triggered_by_item or "factor overall")
            quotes = recs.cell(row=row, column=7,
                               value="\n".join(f"“{q}”" for q in rec.evidence))
            quotes.alignment = Alignment(wrap_text=True, vertical="top")
            for column in range(1, 8):
                recs.cell(row=row, column=column).border = _BORDER
            recs.row_dimensions[row].height = 42
    else:
        recs.cell(row=2, column=1,
                  value="No actions flagged -- every factor is at or above the "
                        "Advanced threshold.")
    recs.freeze_panes = "A2"

    # --- Written answers --------------------------------------------------
    # The three open questions, and what reading them found. Kept on their own
    # sheet so the free text is available in full alongside how it was read.
    if results.notes:
        written = workbook.create_sheet("Written answers")
        _set_widths(written, [(1, 42), (2, 90), (3, 20)])
        _write_header(written, 1, ["Question", "Answer", "Read as being about"])

        insights = results.text_insights
        placed = ", ".join(s.factor_name for s in insights.by_factor) \
            if insights is not None else ""

        row = 2
        for question, answer in results.notes.items():
            written.cell(row=row, column=1, value=question)
            cell = written.cell(row=row, column=2, value=answer)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            written.row_dimensions[row].height = 40
            row += 1

        row += 1
        written.cell(row=row, column=1,
                     value="Factors named in the written answers").font = _TITLE_FONT
        row += 1
        written.cell(row=row, column=1, value=placed or "none identified")
        row += 2
        written.cell(row=row, column=1, value=(
            "These answers are not scored. They are used to attach your own "
            "words to the recommendations, and to raise a factor you wrote "
            "about that the scores did not flag."
        ))
        written.cell(row=row, column=1).font = Font(size=9, color="777777")

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def write_excel(results: AssessmentResults, path: str) -> str:
    """Write the Excel workbook to disk. Returns the path."""
    with open(path, "wb") as handle:
        handle.write(to_excel_bytes(results))
    return path


# ---------------------------------------------------------------------------
# Filenames
# ---------------------------------------------------------------------------

def suggested_filename(results: AssessmentResults, extension: str) -> str:
    """
    A filename for the download button.

    Uses the tier and the date rather than anything identifying, since there's
    no company name collected to use.
    """
    stamp = datetime.now().strftime("%Y-%m-%d")
    tier = results.readiness_tier.lower()
    return f"ai-readiness-{tier}-{stamp}.{extension.lstrip('.')}"


class ExportService:
    """
    Thin object wrapper over the functions above.

    The functions are the real interface; this exists because the class diagram
    names an ExportService, and the dashboard reads more clearly with one.
    """

    def to_pdf(self, results: AssessmentResults) -> bytes:
        return to_pdf_bytes(results)

    def to_excel(self, results: AssessmentResults) -> bytes:
        return to_excel_bytes(results)

    def write_pdf(self, results: AssessmentResults, path: str) -> str:
        return write_pdf(results, path)

    def write_excel(self, results: AssessmentResults, path: str) -> str:
        return write_excel(results, path)

    def filename(self, results: AssessmentResults, extension: str) -> str:
        return suggested_filename(results, extension)


# ---------------------------------------------------------------------------
# Whole-dataset reporting
# ---------------------------------------------------------------------------
# Everything above describes one company. This part describes a set of them,
# and it lives here rather than in the dashboard so that the screen and the
# downloads are computed once, from the same code. When a report disagrees
# with the page that produced it, this is usually why.

from dataclasses import dataclass as _dataclass, field as _field  # noqa: E402


@_dataclass
class GroupProfile:
    """Mean factor scores for one slice of the data (a country, a sector)."""

    label: str
    count: int
    overall: float
    factor_means: List[Tuple[str, float]]   # (factor_id, mean score)


@_dataclass
class DatasetSummary:
    """Everything the dataset report needs, computed once."""

    count: int = 0
    scores: List[float] = _field(default_factory=list)
    mean: float = 0.0
    lowest: float = 0.0
    highest: float = 0.0
    partial: bool = False
    factors_present: List[str] = _field(default_factory=list)
    factors_absent: List[str] = _field(default_factory=list)
    tier_counts: Dict[str, int] = _field(default_factory=dict)
    factor_means: List[Tuple[str, float]] = _field(default_factory=list)
    by_country: List[GroupProfile] = _field(default_factory=list)
    by_sector: List[GroupProfile] = _field(default_factory=list)
    common_barriers: List[Tuple[str, int]] = _field(default_factory=list)
    recommendations: List = _field(default_factory=list)
    recommendation_counts: List[Tuple[str, int]] = _field(default_factory=list)
    themes: List[Tuple[str, int, List[str]]] = _field(default_factory=list)
    notes: List[Tuple[str, Dict[str, str]]] = _field(default_factory=list)
    # What reading the written answers found -- which factors they name, and
    # how many were set aside as non-answers.
    text_insights: Any = None
    source_name: str = ""

    @property
    def split_profiles(self) -> bool:
        """
        Whether the shape should be drawn per group rather than once.

        One averaged shape across several countries or industries hides the
        thing worth seeing. It is only the right picture when the data really
        is one group -- a single company, a single country, or a file that
        never said.
        """
        return len(self.by_country) > 1 or len(self.by_sector) > 1


def _group_profiles(scored, profiles_by_id, attribute: str,
                    minimum: int = 1) -> List[GroupProfile]:
    buckets: Dict[str, List] = {}
    for identifier, result in scored:
        value = getattr(profiles_by_id.get(identifier), attribute, None)
        if value:
            buckets.setdefault(str(value), []).append(result)

    groups: List[GroupProfile] = []
    for label, results in buckets.items():
        if len(results) < minimum:
            continue
        overall = sum(r.overall_score for r in results) / len(results)
        means: List[Tuple[str, float]] = []
        for factor in cfg.FACTORS:
            values = [r.factor(factor.id).score for r in results
                      if factor.id in r.factor_scores]
            if values:
                means.append((factor.id, sum(values) / len(values)))
        groups.append(GroupProfile(label, len(results), overall, means))

    return sorted(groups, key=lambda g: -g.count)


def extract_themes(notes: Sequence[Tuple[str, Dict[str, str]]],
                   limit: int = 7) -> List[Tuple[str, int, List[str]]]:
    """
    Pull recurring subjects out of the written answers.

    The comments are the part of a survey that usually explains the numbers,
    and reading 300 of them by hand is not realistic. This groups them by the
    readiness concepts they mention -- using the same lexicon that matches
    questions -- and reports how many people raised each, with a couple of
    their own sentences as evidence.

    Counting concepts rather than words matters: "cost", "budget", "expensive"
    and "can't afford it" are one concern, and a raw word count would report
    them as four small ones.
    """
    from .survey_import import CONCEPT_LEXICON, _terms

    readable = {
        "money": "Cost and funding",
        "people": "Skills and training",
        "leader": "Leadership and strategy",
        "data": "Data quality and access",
        "govern": "Governance, privacy and trust",
        "culture": "Culture and resistance to change",
        "infra": "Technology and infrastructure",
        "adopt": "Getting started with adoption",
    }

    counts: Dict[str, int] = {}
    candidates: Dict[str, List[str]] = {}

    for _, answers in notes:
        text = " ".join(answers.values())
        mentioned = {
            term[1:] for term in _terms(text)
            if term.startswith("~") and term[1:] in readable
        }
        snippet = " ".join(text.split())
        snippet = snippet[:160] + ("…" if len(snippet) > 160 else "")
        for concept in mentioned:
            counts[concept] = counts.get(concept, 0) + 1
            candidates.setdefault(concept, []).append(snippet)

    ranked = sorted(counts.items(), key=lambda kv: -kv[1])[:limit]

    # One long answer often mentions four subjects at once, and quoting it
    # under all four makes the section read as if the same person said
    # everything. A theme takes a quote nothing else has used where it has one.
    used: set = set()
    themes: List[Tuple[str, int, List[str]]] = []
    for concept, count in ranked:
        pool = candidates.get(concept, [])
        fresh = [quote for quote in pool if quote not in used]
        chosen = (fresh or pool)[:2]
        used.update(chosen)
        themes.append((readable[concept], count, chosen))

    return themes


def cohort_results(scored) -> AssessmentResults:
    """
    Build a single results object representing the average company.

    This exists so the cohort gets recommendations from exactly the same rules
    a single company does, rather than a second set written separately that
    would drift out of step with the first.
    """
    results = [r for _, r in scored]
    factor_scores: Dict[str, FactorScore] = {}

    for factor in cfg.FACTORS:
        present = [r.factor(factor.id) for r in results if factor.id in r.factor_scores]
        if not present:
            continue

        mean_score = sum(f.score for f in present) / len(present)
        items: List[ItemScore] = []
        for item in factor.items:
            values = [
                i.score for f in present for i in f.item_scores if i.item_id == item.id
            ]
            if not values:
                continue
            mean_item = sum(values) / len(values)
            items.append(ItemScore(
                item_id=item.id, text=item.text, factor_id=factor.id,
                factor_name=factor.name,
                raw_value=round(mean_item / 25) + 1,
                adjusted_value=round(mean_item / 25) + 1,
                reverse=item.reverse, score=mean_item,
            ))

        factor_scores[factor.id] = FactorScore(
            factor_id=factor.id, name=factor.name, weight=factor.weight,
            mean_rating=mean_score / 25 + 1, score=mean_score,
            band_label=cfg.band_for_score(mean_score).label,
            weighted_contribution=mean_score * factor.weight,
            item_scores=items, answered=len(items), expected=len(factor.items),
        )

    total_weight = sum(f.weight for f in factor_scores.values()) or 1.0
    overall = sum(f.weighted_contribution for f in factor_scores.values()) / total_weight
    band = cfg.band_for_score(overall)

    strong = sorted((f for f in factor_scores.values()
                     if f.score >= cfg.STRENGTH_THRESHOLD),
                    key=lambda f: -f.score)[:cfg.MAX_STRENGTHS]
    weak = sorted((f for f in factor_scores.values()
                   if f.score < cfg.BARRIER_THRESHOLD),
                  key=lambda f: -f.impact)[:cfg.MAX_BARRIERS]

    gaps: List[ItemScore] = []
    per_factor: Dict[str, int] = {}
    for item in sorted((i for f in factor_scores.values() for i in f.item_scores
                        if i.score < cfg.ITEM_BARRIER_THRESHOLD),
                       key=lambda i: i.score):
        if per_factor.get(item.factor_id, 0) >= cfg.MAX_ITEM_BARRIERS_PER_FACTOR:
            continue
        gaps.append(item)
        per_factor[item.factor_id] = per_factor.get(item.factor_id, 0) + 1
        if len(gaps) >= cfg.MAX_ITEM_BARRIERS:
            break

    return AssessmentResults(
        overall_score=overall, readiness_tier=band.label,
        tier_description=band.description, factor_scores=factor_scores,
        strengths=strong, barriers=weak, item_barriers=gaps,
        partial=any(f.is_partial for f in factor_scores.values())
        or len(factor_scores) < len(cfg.FACTORS),
    )


def summarise_dataset(scored, profiles_by_id,
                      notes: Optional[Sequence[Tuple[str, Dict[str, str]]]] = None,
                      source_name: str = "") -> DatasetSummary:
    """Compute everything the dataset report and the dashboard both need."""
    from .recommendation_engine import RecommendationEngine

    results = [r for _, r in scored]
    overall = [r.overall_score for r in results]
    notes = list(notes or [])

    present = [f.id for f in cfg.FACTORS
               if any(f.id in r.factor_scores for r in results)]

    factor_means: List[Tuple[str, float]] = []
    for factor_id in present:
        values = [r.factor(factor_id).score for r in results
                  if factor_id in r.factor_scores]
        factor_means.append((factor_id, sum(values) / len(values)))

    barrier_counts: Dict[str, int] = {}
    rec_counts: Dict[str, int] = {}
    for result in results:
        for factor in result.barriers:
            barrier_counts[factor.name] = barrier_counts.get(factor.name, 0) + 1
        for rec in result.recommendations:
            rec_counts[rec.title] = rec_counts.get(rec.title, 0) + 1

    from .text_analysis import analyse_notes, apply_to_recommendations

    average = cohort_results(scored)
    RecommendationEngine().recommend(average)

    # The cohort's advice comes from the scored average, then the whole
    # dataset's written answers are read against it: the quotes attach to the
    # actions they support, and a factor many people write about but no score
    # flagged is added as its own, separately labelled item.
    insights = analyse_notes(notes)
    apply_to_recommendations(average, insights)

    return DatasetSummary(
        count=len(results),
        scores=list(overall),
        mean=sum(overall) / len(overall) if overall else 0.0,
        lowest=min(overall) if overall else 0.0,
        highest=max(overall) if overall else 0.0,
        partial=any(r.partial for r in results),
        factors_present=present,
        factors_absent=[f.id for f in cfg.FACTORS if f.id not in present],
        tier_counts={
            band.label: sum(1 for r in results if r.readiness_tier == band.label)
            for band in cfg.READINESS_BANDS
        },
        factor_means=sorted(factor_means, key=lambda kv: -kv[1]),
        by_country=_group_profiles(scored, profiles_by_id, "region"),
        by_sector=_group_profiles(scored, profiles_by_id, "industry_sector"),
        common_barriers=sorted(barrier_counts.items(), key=lambda kv: -kv[1]),
        recommendations=average.recommendations,
        recommendation_counts=sorted(rec_counts.items(), key=lambda kv: -kv[1])[:10],
        themes=extract_themes(notes),
        notes=notes,
        text_insights=insights,
        source_name=source_name,
    )


def dataset_pdf_bytes(summary: DatasetSummary) -> bytes:
    """The whole-dataset report, as a PDF."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4,
        leftMargin=18 * mm, rightMargin=18 * mm,
        topMargin=16 * mm, bottomMargin=16 * mm,
        title="AI Adoption Readiness — dataset analysis",
    )
    s = _styles()
    width = doc.width
    story: List = []

    story.append(Paragraph("AI Adoption Readiness", s["title"]))
    story.append(Paragraph(
        f"Dataset analysis &middot; {summary.count} responses &middot; "
        f"generated {_timestamp()}", s["subtitle"]))

    band_colour = BAND_COLOURS.get(cfg.band_for_score(summary.mean).label, "555555")
    headline = Table([[
        Paragraph(f'<font size="30"><b>{summary.mean:.1f}</b></font>'
                  f'<font size="12" color="#777777">/100</font>', s["body"]),
        Paragraph(
            f'<b>Mean readiness across {summary.count} companies</b><br/>'
            f'Lowest {summary.lowest:.1f} &nbsp;&middot;&nbsp; '
            f'Highest {summary.highest:.1f}<br/>'
            + "&nbsp;&middot;&nbsp; ".join(
                f"{label} {count}" for label, count in summary.tier_counts.items()),
            s["body"]),
    ]], colWidths=[38 * mm, width - 38 * mm])
    headline.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F7F7F7")),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ("LEFTPADDING", (0, 0), (-1, -1), 10),
        ("LINEBEFORE", (0, 0), (0, 0), 3, colors.HexColor("#" + band_colour)),
    ]))
    story.append(headline)

    if summary.partial:
        story.append(Spacer(1, 8))
        absent = ", ".join(cfg.FACTORS_BY_ID[f].name for f in summary.factors_absent)
        story.append(Paragraph(
            "<b>Partial coverage.</b> This dataset does not answer every "
            "question, so the scores below are built from the factors it does "
            "cover and re-weighted across them."
            + (f" No reading was possible for: {escape(absent)}." if absent else ""),
            s["small"]))

    # --- how to read it ---------------------------------------------------
    # At the top, not the back. Every number below is on a scale the reader has
    # no reason to already know, and an explanation placed after the figures it
    # explains has been read too late to be any use.
    story.append(Paragraph("How to read this", s["h2"]))
    story.append(Paragraph(
        "Each factor is scored from its questions on a 1&ndash;5 agreement "
        "scale converted to 0&ndash;100, with barrier-worded questions "
        "reversed first, so a higher score always means more ready. Group "
        "figures are means, and a group of one company is that company rather "
        "than an average. Bands: Emerging 0&ndash;40, Developing 40&ndash;70, "
        "Advanced 70&ndash;100.", s["small"]))

    # --- factor averages --------------------------------------------------
    story.append(Paragraph("Average score by factor", s["h2"]))
    rows: List[List] = []
    for factor_id, mean in summary.factor_means:
        factor = cfg.FACTORS_BY_ID[factor_id]
        band = cfg.band_for_score(mean).label
        rows.append([
            Paragraph(escape(factor.name), s["body"]),
            Paragraph(f"<b>{mean:.1f}</b>", s["body"]),
            _score_bar(mean, width - 100 * mm, BAND_COLOURS.get(band, "555555")),
            Paragraph(f'<font color="#{BAND_COLOURS.get(band, "555555")}">'
                      f'{band}</font>', s["small"]),
        ])
    table = Table(rows, colWidths=[62 * mm, 16 * mm, width - 100 * mm, 22 * mm])
    table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (0, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LINEBELOW", (0, 0), (-1, -2), 0.4, colors.HexColor("#EAEAEA")),
    ]))
    story.append(table)

    # --- profile shapes ---------------------------------------------------
    # One shape per country and per sector when the data holds more than one,
    # because a single averaged shape across several of them shows a company
    # that does not exist.
    def shapes_for(groups: List[GroupProfile], heading: str) -> None:
        if not groups:
            return
        story.append(Paragraph(heading, s["h2"]))
        cells, labels = [], []
        for group in groups[:6]:
            scores = [(SHORT_FACTOR_LABELS.get(fid, fid), value)
                      for fid, value in group.factor_means]
            colour = BAND_COLOURS.get(cfg.band_for_score(group.overall).label, "555555")
            cells.append(_radar_drawing(scores, colour, size=118))
            labels.append(Paragraph(
                f"<b>{escape(group.label)}</b><br/>"
                f"{group.overall:.1f}/100 &middot; n={group.count}", s["small"]))

        for start in range(0, len(cells), 3):
            chunk = cells[start:start + 3]
            chunk_labels = labels[start:start + 3]
            while len(chunk) < 3:
                chunk.append("")
                chunk_labels.append("")
            grid = Table([chunk, chunk_labels], colWidths=[width / 3.0] * 3)
            grid.setStyle(TableStyle([
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("BOTTOMPADDING", (0, 1), (-1, 1), 10),
            ]))
            story.append(grid)

    if summary.split_profiles:
        shapes_for(summary.by_country, "Profile shape by country")
        shapes_for(summary.by_sector, "Profile shape by sector")
    else:
        story.append(Paragraph("Profile shape", s["h2"]))
        scores = [(SHORT_FACTOR_LABELS.get(fid, fid), value)
                  for fid, value in summary.factor_means]
        story.append(_radar_drawing(scores, band_colour, size=175))

    # --- recommendations --------------------------------------------------
    story.append(Paragraph("Recommended actions for this group", s["h2"]))
    if not summary.recommendations:
        story.append(Paragraph(
            "No actions flagged — every factor averages at or above the "
            "Advanced threshold.", s["body"]))
    else:
        story.append(Paragraph(
            "Generated from the average profile across the dataset, using the "
            "same rules applied to an individual assessment.", s["small"]))
        story.append(Spacer(1, 6))
        for index, rec in enumerate(summary.recommendations, start=1):
            colour = SEVERITY_COLOURS.get(rec.severity, "555555")
            label = SEVERITY_LABELS.get(rec.severity, rec.severity.title())
            block = Table([[Paragraph(
                f'<font color="#{colour}"><b>{index}. {escape(label)}</b></font>'
                f'&nbsp;&nbsp;<font color="#777777" size="8">'
                f'{escape(rec.factor_name)}</font><br/>'
                f'<b>{escape(rec.title)}</b><br/>{escape(rec.action)}'
                f'{_evidence_markup(rec)}', s["body"])]],
                colWidths=[width])
            block.setStyle(TableStyle([
                ("LINEBEFORE", (0, 0), (0, 0), 2.5, colors.HexColor("#" + colour)),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FAFAFA")),
            ]))
            story.append(KeepTogether(block))
            story.append(Spacer(1, 5))

    # --- what people wrote ------------------------------------------------
    insights = summary.text_insights
    if insights is not None and insights.used:
        story.append(Paragraph("Which factors the written answers name", s["h2"]))
        story.append(Paragraph(
            f"{insights.answers_used} of {insights.answers_seen} written "
            f"answers said something specific; the remaining "
            f"{insights.answers_dropped} were blanks or non-answers "
            f"(&ldquo;N/A&rdquo;, &ldquo;none&rdquo;, &ldquo;no&rdquo;) and "
            f"were set aside rather than counted. Each remaining answer is "
            f"placed against the readiness factor it talks about, so the "
            f"comments can be compared with the scores. Shares below are of "
            f"all {insights.total_respondents} respondents, not only those "
            f"whose answers could be read. They do not change any score.",
            s["small"]))
        story.append(Spacer(1, 5))

        rows = [[
            Paragraph("<b>Factor</b>", s["small"]),
            Paragraph("<b>Raised by</b>", s["small"]),
            Paragraph("<b>Mean score</b>", s["small"]),
        ]]
        means = dict(summary.factor_means)
        for signal in insights.by_factor:
            score = means.get(signal.factor_id)
            rows.append([
                Paragraph(escape(signal.factor_name), s["small"]),
                Paragraph(f"{signal.mentions} of {insights.total_respondents} "
                          f"({signal.share:.0f}%)", s["small"]),
                Paragraph(f"{score:.1f}/100" if score is not None
                          else "not scored", s["small"]),
            ])
        placement = Table(rows, colWidths=[width - 76 * mm, 42 * mm, 34 * mm])
        placement.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (0, -1), 0),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("LINEBELOW", (0, 0), (-1, -2), 0.4, colors.HexColor("#EAEAEA")),
        ]))
        story.append(placement)

    if summary.themes:
        story.append(Paragraph("What respondents raised themselves", s["h2"]))
        story.append(Paragraph(
            f"Grouped from the written answers in {len(summary.notes)} "
            f"responses, by subject rather than by word, so &ldquo;cost&rdquo;, "
            f"&ldquo;budget&rdquo; and &ldquo;can't afford it&rdquo; count as "
            f"one concern.", s["small"]))
        story.append(Spacer(1, 5))
        for name, count, examples in summary.themes:
            share = count / max(1, len(summary.notes)) * 100
            story.append(Paragraph(
                f"<b>{escape(name)}</b> &mdash; raised by {count} "
                f"({share:.0f}%)", s["body"]))
            for example in examples:
                story.append(Paragraph(
                    f'<i>&ldquo;{escape(example)}&rdquo;</i>', s["small"]))
            story.append(Spacer(1, 5))

    doc.build(story)
    return buffer.getvalue()


def dataset_excel_bytes(summary: DatasetSummary) -> bytes:
    """The whole-dataset analysis, as a workbook with charts."""
    workbook = Workbook()

    overview = workbook.active
    overview.title = "Overview"
    _set_widths(overview, [(1, 34), (2, 22), (3, 14)])
    overview["A1"] = "AI Adoption Readiness — dataset analysis"
    overview["A1"].font = _TITLE_FONT
    overview["A2"] = f"Generated {_timestamp()}"
    overview["A2"].font = Font(size=9, color="777777")

    row = 4
    for label, value in (
        ("Responses scored", summary.count),
        ("Mean readiness", round(summary.mean, 1)),
        ("Lowest", round(summary.lowest, 1)),
        ("Highest", round(summary.highest, 1)),
        ("Coverage", "partial" if summary.partial else "all 32 questions"),
        ("Source file", summary.source_name or "—"),
    ):
        overview.cell(row=row, column=1, value=label).font = Font(bold=True)
        overview.cell(row=row, column=2, value=value)
        row += 1

    row += 1
    overview.cell(row=row, column=1, value="Tier split").font = _TITLE_FONT
    row += 1
    tier_start = row
    _write_header(overview, row, ["Tier", "Companies", "Share %"])
    row += 1
    for label, count in summary.tier_counts.items():
        overview.cell(row=row, column=1, value=label)
        overview.cell(row=row, column=2, value=count)
        overview.cell(row=row, column=3,
                      value=round(count / max(1, summary.count) * 100, 1))
        row += 1

    tiers = BarChart()
    tiers.type = "col"
    tiers.title = "Companies by readiness tier"
    tiers.add_data(Reference(overview, min_col=2, min_row=tier_start,
                             max_row=row - 1), titles_from_data=True)
    _label_chart(tiers,
                 Reference(overview, min_col=1, min_row=tier_start + 1,
                           max_row=row - 1),
                 category_title="Readiness tier",
                 value_title="Number of companies")
    tiers.height, tiers.width, tiers.legend = 8, 14, None
    overview.add_chart(tiers, f"E{tier_start}")

    # --- factors ----------------------------------------------------------
    factors = workbook.create_sheet("Factor averages")
    _set_widths(factors, [(1, 42), (2, 14), (3, 14), (4, 14)])
    # The short label is a real column rather than something the chart invents,
    # so the reader can see which abbreviation belongs to which factor.
    _write_header(factors, 1, ["Factor", "Mean score", "Band", "Short label"])
    for index, (factor_id, mean) in enumerate(summary.factor_means, start=2):
        band = cfg.band_for_score(mean).label
        factors.cell(row=index, column=1, value=cfg.FACTORS_BY_ID[factor_id].name)
        factors.cell(row=index, column=2, value=round(mean, 1))
        cell = factors.cell(row=index, column=3, value=band)
        cell.font = Font(bold=True, color=BAND_COLOURS.get(band, "000000"))
        factors.cell(row=index, column=4,
                     value=SHORT_FACTOR_LABELS.get(factor_id, factor_id))
        for column in range(1, 5):
            factors.cell(row=index, column=column).border = _BORDER

    last = len(summary.factor_means) + 1
    labels = Reference(factors, min_col=1, min_row=2, max_row=last)
    values = Reference(factors, min_col=2, min_row=1, max_row=last)
    # Full names round a radar overlap into each other; the bar chart has the
    # room for them and the radar does not.
    short_labels = Reference(factors, min_col=4, min_row=2, max_row=last)

    bars = BarChart()
    bars.type = "bar"
    bars.title = "Average score by factor (0-100)"
    bars.add_data(values, titles_from_data=True)
    _label_chart(bars, labels,
                 category_title="Readiness factor",
                 value_title="Mean score /100", value_range=(0, 100))
    bars.height, bars.width, bars.legend = 9, 20, None
    factors.add_chart(bars, f"A{last + 3}")

    shape = RadarChart()
    shape.type = "filled"
    shape.title = "Average profile shape (0-100 by factor)"
    shape.add_data(values, titles_from_data=True)
    _label_chart(shape, short_labels, show_values=False, value_axis=False)
    shape.y_axis.scaling.min, shape.y_axis.scaling.max = 0, 100
    shape.height, shape.width = 11, 11
    factors.add_chart(shape, f"H{last + 3}")

    # --- groups -----------------------------------------------------------
    for title, groups in (("By country", summary.by_country),
                          ("By sector", summary.by_sector)):
        if not groups:
            continue
        sheet = workbook.create_sheet(title)
        headers = [title.replace("By ", "").title(), "Companies", "Mean"]
        headers += [cfg.FACTORS_BY_ID[f.id].name for f in cfg.FACTORS]
        _set_widths(sheet, [(1, 26), (2, 12), (3, 10)]
                    + [(i, 20) for i in range(4, 11)])
        _write_header(sheet, 1, headers)

        for index, group in enumerate(groups, start=2):
            means = dict(group.factor_means)
            sheet.cell(row=index, column=1, value=group.label)
            sheet.cell(row=index, column=2, value=group.count)
            sheet.cell(row=index, column=3, value=round(group.overall, 1))
            for offset, factor in enumerate(cfg.FACTORS):
                value = means.get(factor.id)
                sheet.cell(row=index, column=4 + offset,
                           value=round(value, 1) if value is not None else None)

        end = len(groups) + 1
        sheet.freeze_panes = "A2"

        # A second, smaller block holding just the factor scores, which is what
        # the radar is drawn from. A radar needs the factor columns directly
        # beside the group name and the table above has the count and the mean
        # in between, so the numbers are laid out once more in the shape the
        # chart can read. Short labels, because full factor names overlap each
        # other round a radar.
        present = [factor_id for factor_id, _ in summary.factor_means]
        shown = min(len(groups), 6)
        split = bool(present) and len(groups) > 1

        chart_row = end + 3
        head = 0
        if split:
            sheet.cell(row=end + 2, column=1,
                       value=f"Profile shape data ({title.lower()})").font = _TITLE_FONT
            head = end + 3
            _write_header(sheet, head, ["Group"] + [
                SHORT_FACTOR_LABELS.get(f, f) for f in present])
            for offset, group in enumerate(groups[:6], start=1):
                means = dict(group.factor_means)
                sheet.cell(row=head + offset, column=1, value=group.label)
                for column, factor_id in enumerate(present, start=2):
                    value = means.get(factor_id)
                    sheet.cell(row=head + offset, column=column,
                               value=round(value, 1) if value is not None else 0)
            chart_row = head + shown + 2

        chart = BarChart()
        chart.type = "col"
        chart.title = f"Mean readiness {title.lower()}"
        chart.add_data(Reference(sheet, min_col=3, min_row=1, max_row=end),
                       titles_from_data=True)
        _label_chart(chart,
                     Reference(sheet, min_col=1, min_row=2, max_row=end),
                     category_title=headers[0],
                     value_title="Mean score /100", value_range=(0, 100))
        chart.height, chart.width, chart.legend = 8, 18, None
        sheet.add_chart(chart, f"A{chart_row}")

        # One radar carrying a series per group: the workbook's version of the
        # split profile shapes in the PDF. A single averaged shape across
        # several countries or sectors describes a company that does not exist.
        if split:
            shape = RadarChart()
            shape.type = "marker"
            shape.title = f"Profile shape {title.lower()} (0-100)"
            shape.add_data(
                Reference(sheet, min_col=1, max_col=len(present) + 1,
                          min_row=head + 1, max_row=head + shown),
                titles_from_data=True, from_rows=True,
            )
            _label_chart(
                shape,
                Reference(sheet, min_col=2, max_col=len(present) + 1,
                          min_row=head, max_row=head),
                show_values=False, value_axis=False,
            )
            shape.y_axis.scaling.min, shape.y_axis.scaling.max = 0, 100
            shape.height, shape.width = 12, 15
            sheet.add_chart(shape, f"L{chart_row}")

    # --- recommendations and themes ---------------------------------------
    advice = workbook.create_sheet("Recommendations")
    _set_widths(advice, [(1, 9), (2, 18), (3, 32), (4, 40), (5, 70), (6, 60)])
    _write_header(advice, 1, ["Priority", "Severity", "Factor", "Action",
                              "Detail", "Supporting words from respondents"])
    for index, rec in enumerate(summary.recommendations, start=1):
        row = index + 1
        advice.cell(row=row, column=1, value=index)
        cell = advice.cell(row=row, column=2,
                           value=SEVERITY_LABELS.get(rec.severity, rec.severity))
        cell.font = Font(bold=True, color=SEVERITY_COLOURS.get(rec.severity, "000000"))
        advice.cell(row=row, column=3, value=rec.factor_name)
        advice.cell(row=row, column=4, value=rec.title)
        detail = advice.cell(row=row, column=5, value=rec.action)
        detail.alignment = Alignment(wrap_text=True, vertical="top")
        quotes = advice.cell(row=row, column=6,
                             value="\n".join(f"“{q}”" for q in rec.evidence))
        quotes.alignment = Alignment(wrap_text=True, vertical="top")
        advice.row_dimensions[row].height = 42

    insights = summary.text_insights
    if summary.themes or (insights is not None and insights.used):
        themes = workbook.create_sheet("Written answers")
        _set_widths(themes, [(1, 42), (2, 14), (3, 90)])
        row = 1

        if insights is not None and insights.used:
            themes.cell(row=row, column=1,
                        value="How the written answers were read").font = _TITLE_FONT
            row += 1
            for label, value in (
                ("Respondents with written answers", insights.total_respondents),
                ("Answers looked at", insights.answers_seen),
                ("Answers used", insights.answers_used),
                ("Set aside as non-answers", insights.answers_dropped),
                ("Respondents who wrote something usable", insights.respondents),
            ):
                themes.cell(row=row, column=1, value=label).font = Font(bold=True)
                themes.cell(row=row, column=2, value=value)
                row += 1
            if insights.drop_reasons:
                themes.cell(row=row, column=1, value="Why answers were set aside") \
                    .font = Font(bold=True)
                themes.cell(row=row, column=2, value=", ".join(
                    f"{reason}: {count}"
                    for reason, count in sorted(insights.drop_reasons.items(),
                                                key=lambda kv: -kv[1])))
                row += 1

            row += 1
            themes.cell(row=row, column=1,
                        value="Factors named in the written answers").font = _TITLE_FONT
            row += 1
            _write_header(themes, row, ["Factor", "Raised by", "Example"])
            row += 1
            for signal in insights.by_factor:
                themes.cell(row=row, column=1, value=signal.factor_name)
                themes.cell(row=row, column=2,
                            value=f"{signal.mentions} ({signal.share:.0f}%)")
                cell = themes.cell(row=row, column=3,
                                   value=signal.examples[0] if signal.examples else "")
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                row += 1
            row += 1

        themes.cell(row=row, column=1, value="Recurring subjects").font = _TITLE_FONT
        row += 1
        _write_header(themes, row, ["Theme", "Mentions", "Example"])
        row += 1
        for name, count, examples in summary.themes:
            themes.cell(row=row, column=1, value=name)
            themes.cell(row=row, column=2, value=count)
            example = themes.cell(row=row, column=3,
                                  value=examples[0] if examples else "")
            example.alignment = Alignment(wrap_text=True, vertical="top")
            row += 1

        row += 1
        themes.cell(row=row, column=1, value="All written answers").font = _TITLE_FONT
        row += 1
        _write_header(themes, row, ["Response", "Question", "Answer"])
        row += 1
        for identifier, answers in summary.notes:
            for question, answer in answers.items():
                themes.cell(row=row, column=1, value=identifier)
                themes.cell(row=row, column=2, value=question)
                cell = themes.cell(row=row, column=3, value=answer)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                row += 1

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
